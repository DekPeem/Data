"""ตัวอ่านไฟล์ export จากระบบ "AMI" ของ PEA (โครงการติดตั้งระบบมิเตอร์อัจฉริยะ) — ต่างจาก
pea_ingest.py (ไฟล์ HTML แฝงเป็น .xls จากเว็บ AMRWEB) เพราะไฟล์นี้เป็น Excel (.xlsx) แท้ๆ
(ZIP/OOXML) ตรวจพบครั้งแรกจากไฟล์จริงที่ผู้ใช้ส่งมา (ชื่อไฟล์ขึ้นต้นด้วย "kW" หัวรายงานระบุ
"โครงการติดตั้งระบบมิเตอร์อัจฉริยะ (AMI)")

รองรับ 2 รูปแบบตาราง (ทั้งคู่หัวรายงานขึ้นต้นด้วย "โครงการติดตั้งระบบมิเตอร์อัจฉริยะ (AMI)"
เหมือนกัน แยกกันแค่โครงสร้างตารางข้อมูลราย 15 นาที):

  1. "ข้อมูลกิโลวัตต์รายเดือน" — คอลัมน์ Rate A/B/C ตามรอบ TOU เดียวกับ pea_ingest (ดู
     pea_ingest.RATE_TO_PERIOD) ป้ายชื่อหัวรายงานมีคำว่า "ไฟฟ้า" ต่อท้าย (เช่น "บัญชีผู้ใช้ไฟฟ้า"
     แทน "บัญชีผู้ใช้ไฟ") และปีในคอลัมน์เวลาเป็นปี พ.ศ. (ต้องแปลงเป็น ค.ศ. ก่อน)

  2. "Custom kW Report" — ไม่มีคอลัมน์ Rate A/B/C เลย มีแค่ [Time, kW] เท่านั้น (ต้องคำนวณช่วง
     P/OP/H เองจาก timestamp ผ่าน pea_ingest.classify_tou_period แล้วแปลง kW เป็น kWh ของช่วง
     15 นาที คือ kW x 0.25 — เหมือนหลักการเดียวกับ pea_meter_log_ingest.py) ป้ายชื่อหัวรายงาน
     เป็นภาษาอังกฤษ ("Contact Account :", "Meter No. :") และปีในคอลัมน์เวลาเป็น ค.ศ. อยู่แล้ว
     (ไม่ต้องแปลง พ.ศ.)

⚠️ หลักการความปลอดภัยเดียวกับ pea_ingest.py: ไฟล์ดิบมีข้อมูลระบุตัวตนลูกค้า อ่านเฉพาะตัวเลข
การใช้ไฟฟ้าเพื่อคำนวณค่าเฉลี่ยแบบ anonymized เท่านั้น ไม่ควร commit ไฟล์ดิบเข้า repo public
"""

from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path
from typing import List, Union

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise ImportError(
        "ต้องติดตั้ง openpyxl ก่อนใช้งาน pea_ami_ingest: pip install openpyxl"
    ) from exc

from .pea_ingest import IntervalReading, _to_float, classify_tou_period, is_ami_xlsx  # noqa: F401 (re-exported for callers)

# ป้ายชื่อหัวรายงานที่ต้องการอ่าน — คีย์ฝั่งซ้ายคือคำที่ปรากฏจริงในไฟล์ AMI (รูปแบบ "ข้อมูลกิโลวัตต์
# รายเดือน" มี "ไฟฟ้า" ต่อท้าย, รูปแบบ "Custom kW Report" เป็นภาษาอังกฤษ) คีย์ฝั่งขวาคือชื่อ
# มาตรฐานเดียวกับที่ pea_ingest.parse_report_header ใช้ แปลงให้ตรงกันเพื่อให้โค้ดฝั่งเรียกใช้
# (amr_import.py) ใช้ key เดียวกันได้ไม่ต้องรู้ว่าไฟล์เป็นรูปแบบไหน
_LABEL_ALIASES = {
    "บัญชีผู้ใช้ไฟฟ้า": "บัญชีผู้ใช้ไฟ",
    "ชื่อผู้ใช้ไฟฟ้า": "ชื่อผู้ใช้ไฟ",
    "Contact Account": "บัญชีผู้ใช้ไฟ",
    "Meter No.": "หมายเลขมิเตอร์",
}
_KNOWN_LABELS = {"บัญชีผู้ใช้ไฟ", "ชื่อผู้ใช้ไฟ", "หมายเลขมิเตอร์", "Tariff", "CT Ratio", "VT Ratio"}

_TIMESTAMP_RE = re.compile(r"^(\d{2})/(\d{2})/(\d{4})(\s.*)$")


# ปีที่มากกว่านี้ถือว่าเป็น พ.ศ. แน่นอน (ไฟล์ AMR ทั้งหมดเป็นข้อมูลปัจจุบัน ไม่มีทางเป็นปี ค.ศ.
# เกิน 2100 ไปได้จริงๆ) — ไฟล์ AMI ที่เจอมาไม่ได้ใช้ พ.ศ. เสมอไป (บางไฟล์ใช้ ค.ศ. อยู่แล้วเหมือน
# ไฟล์ AMRWEB) เลยต้องเช็คก่อนแปลง ไม่ใช่ลบ 543 ทื่อๆ ทุกไฟล์
_BE_YEAR_THRESHOLD = 2100


def _convert_be_to_ce_timestamp(timestamp: str) -> str:
    """แปลงปี พ.ศ. เป็น ค.ศ. ในสตริง timestamp ดิบ ถ้าเป็น พ.ศ. จริง (เช่น "01/09/2568 00.15"
    -> "01/09/2025 00.15") — ไฟล์ AMI บางไฟล์ใช้ พ.ศ. บางไฟล์ใช้ ค.ศ. อยู่แล้วไม่แน่นอน จึงต้อง
    เช็คตัวเลขปีก่อนแปลง (ดู _BE_YEAR_THRESHOLD) ถ้า parse ไม่ได้ หรือเป็น ค.ศ. อยู่แล้ว
    คืนค่าเดิมโดยไม่แตะต้อง"""

    m = _TIMESTAMP_RE.match(timestamp.strip())
    if not m:
        return timestamp
    day, month, year, rest = m.groups()
    year_int = int(year)
    if year_int <= _BE_YEAR_THRESHOLD:
        return timestamp
    return f"{day}/{month}/{year_int - 543}{rest}"


def parse_ami_report_header(path: Union[str, Path]) -> dict:
    """อ่านหัวรายงาน (บัญชีผู้ใช้ไฟ/ชื่อผู้ใช้ไฟ/หมายเลขมิเตอร์/Tariff/CT-VT Ratio) จากไฟล์
    AMI .xlsx — คืน dict คีย์แบบเดียวกับ pea_ingest.parse_report_header (ดู _LABEL_ALIASES)
    เพื่อให้ผู้เรียกใช้ key เดียวกันได้โดยไม่ต้องรู้ว่าไฟล์เป็นรูปแบบไหน สแกนทุกเซลล์ของทุก
    sheet ในไฟล์ (ไม่ใช่แค่ sheet แรก) แทนการอิงตำแหน่งแถว/คอลัมน์/sheet ตายตัว — พบว่าไฟล์ AMI
    บางไฟล์แบ่งหัวรายงาน/ตารางข้อมูลราย 15 นาที/สรุปท้ายตาราง ไว้คนละ sheet กัน (Sheet1/2/3)
    ไม่ได้รวมอยู่ใน sheet เดียวเสมอไปแบบที่เจอไฟล์แรก"""

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        info: dict = {}
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=30):
                for i, cell in enumerate(row):
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    cleaned = value.strip().rstrip(":").strip()
                    label = _LABEL_ALIASES.get(cleaned, cleaned)
                    if label not in _KNOWN_LABELS or label in info:
                        continue
                    next_cell = row[i + 1] if i + 1 < len(row) else None
                    next_value = next_cell.value if next_cell is not None else None
                    info[label] = str(next_value).strip() if next_value is not None else ""
        return info
    finally:
        wb.close()


def _find_rate_columns(ws) -> tuple:
    """หาแถวหัวตาราง Rate A/B/C ใน sheet ที่ระบุ — คืน (เลขแถว, {period: col_idx}) หรือ
    (None, {}) ถ้าไม่เจอใน sheet นี้ (ให้ผู้เรียกลอง sheet อื่นต่อ) ถ้า Rate หนึ่งมีมากกว่า 1
    คอลัมน์ (เจอในบางไฟล์ — คอลัมน์ซ้ำค่าเดียวกัน) ใช้คอลัมน์สุดท้ายที่เจอ ซึ่งยังมีค่าอยู่เสมอ

    นับเลขแถวเองด้วย enumerate (ไม่ใช้ cell.row) เพราะโหมด read_only ของ openpyxl คืน
    EmptyCell ให้คอลัมน์แรกของแถวถ้าคอลัมน์นั้นไม่เคยมีค่าเลย ซึ่งไม่มี .row attribute"""

    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=40), start=1):
        cells_upper = [(str(c.value).strip().upper() if c.value is not None else "") for c in row]
        if sum(1 for c in cells_upper if "RATE" in c) < 2:
            continue
        rate_col_idx: dict = {}
        for i, c in enumerate(cells_upper):
            for suffix, period in (("A", "P"), ("B", "OP"), ("C", "H")):
                if c == f"RATE {suffix}":
                    rate_col_idx[period] = i
        if len(rate_col_idx) >= 2:
            return row_idx, rate_col_idx
    return None, {}


def _parse_ddmmyyyy_hm(timestamp: str):
    """แปลง "DD/MM/YYYY HH.MM" (ค.ศ. แล้ว) เป็น datetime — บางไฟล์ "Custom kW Report" ใช้
    "24.00" แทนเที่ยงคืนของ "วันถัดไป" ไม่ใช่ 00:00 ของวันเดียวกัน (พบจากไฟล์จริงที่ผู้ใช้ส่งมา)
    ซึ่ง datetime.strptime ปกติ parse ไม่ได้ (ชั่วโมงเกิน 23) — เลื่อนไปเป็นวันถัดไป 00:00 แทนเสมอ
    คืน None ถ้า parse ไม่ได้จริงๆ (กันไฟล์แปลกๆ ไม่ให้ทำให้ทั้งการคำนวณพัง)"""

    try:
        return _dt.datetime.strptime(timestamp, "%d/%m/%Y %H.%M")
    except ValueError:
        if timestamp.endswith(" 24.00"):
            try:
                base = _dt.datetime.strptime(timestamp[: -len(" 24.00")], "%d/%m/%Y")
            except ValueError:
                return None
            return base + _dt.timedelta(days=1)
        return None


def _find_time_kw_columns(ws) -> tuple:
    """หาแถวหัวตาราง "Time"/"kW" ใน sheet ที่ระบุ (รายงานแบบ "Custom kW Report" ที่ไม่มีคอลัมน์
    Rate A/B/C ให้เลย — ดู docstring ของโมดูล) คืน (เลขแถว, คอลัมน์ Time, คอลัมน์ kW) หรือ
    (None, None, None) ถ้าไม่เจอใน sheet นี้ (ให้ผู้เรียกลอง sheet อื่นต่อ)"""

    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=40), start=1):
        cells_upper = [(str(c.value).strip().upper() if c.value is not None else "") for c in row]
        if "TIME" in cells_upper and "KW" in cells_upper:
            return row_idx, cells_upper.index("TIME"), cells_upper.index("KW")
    return None, None, None


def parse_ami_interval_report(path: Union[str, Path]) -> List[IntervalReading]:
    """อ่านตารางข้อมูลราย 15 นาทีจากไฟล์ AMI .xlsx — คอลัมน์ Excel จริงแทน <td> ของ HTML
    รองรับ 2 โครงสร้างตาราง (ดู docstring ของโมดูล):

      1. Rate A/B/C — ปีในคอลัมน์เวลาเป็น พ.ศ. (แปลงเป็น ค.ศ. ก่อนคืนค่า — ดู
         _convert_be_to_ce_timestamp)
      2. Time/kW (ไม่มี Rate columns) — คำนวณช่วง P/OP/H เองจาก timestamp ผ่าน
         pea_ingest.classify_tou_period แล้วแปลง kW เป็น kWh ของช่วง 15 นาที (kW x 0.25)

    ค้นหาตารางในทุก sheet ของไฟล์ (ไม่ใช่แค่ sheet แรก) — พบไฟล์จริงที่แบ่งหัวรายงานไว้ sheet
    หนึ่ง (Sheet1) แต่ตารางข้อมูลราย 15 นาทีอยู่อีก sheet หนึ่ง (Sheet2) ลองหาตาราง Rate A/B/C
    ก่อนเสมอ ถ้าไม่เจอเลยสักคอลัมน์ในทุก sheet ถึงจะลองหาตาราง Time/kW แทน"""

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        target_ws = None
        header_row_idx = None
        rate_col_idx: dict = {}
        for ws in wb.worksheets:
            header_row_idx, rate_col_idx = _find_rate_columns(ws)
            if header_row_idx is not None:
                target_ws = ws
                break

        if target_ws is not None:
            readings: List[IntervalReading] = []
            for row in target_ws.iter_rows(min_row=header_row_idx + 1):
                timestamp_cell = row[0].value if len(row) else None
                timestamp = str(timestamp_cell).strip() if timestamp_cell is not None else ""
                if not _TIMESTAMP_RE.match(timestamp):
                    continue  # ข้ามแถวสรุปท้ายตาราง (กิโลวัตต์สูงสุด/คำอธิบาย/พิมพ์โดย ฯลฯ)
                timestamp = _convert_be_to_ce_timestamp(timestamp)
                for period, col_idx in rate_col_idx.items():
                    if col_idx >= len(row):
                        continue
                    raw = row[col_idx].value
                    val = _to_float(str(raw)) if raw is not None else None
                    if val is not None:
                        readings.append(IntervalReading(timestamp=timestamp, period=period, kwh=val))
            return readings

        time_col = kw_col = None
        for ws in wb.worksheets:
            header_row_idx, time_col, kw_col = _find_time_kw_columns(ws)
            if header_row_idx is not None:
                target_ws = ws
                break

        if target_ws is None:
            raise ValueError(
                f'ไม่พบตารางรายงานราย 15 นาที (header ต้องมีคอลัมน์ Rate A/B/C หรือ "Time"/"kW") ในไฟล์ {path}'
            )

        readings = []
        for row in target_ws.iter_rows(min_row=header_row_idx + 1):
            timestamp_cell = row[time_col].value if time_col < len(row) else None
            timestamp = str(timestamp_cell).strip() if timestamp_cell is not None else ""
            if not _TIMESTAMP_RE.match(timestamp):
                continue
            timestamp = _convert_be_to_ce_timestamp(timestamp)
            dt = _parse_ddmmyyyy_hm(timestamp)
            if dt is None:
                continue
            timestamp = dt.strftime("%d/%m/%Y %H.%M")

            raw = row[kw_col].value if kw_col < len(row) else None
            val = _to_float(str(raw)) if raw is not None else None
            if val is None:
                continue
            readings.append(IntervalReading(timestamp=timestamp, period=classify_tou_period(dt), kwh=round(val * 0.25, 4)))
        return readings
    finally:
        wb.close()
