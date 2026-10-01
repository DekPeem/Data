import datetime as dt
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pytest

from amr_mapping.amr_boxplot import (
    ParsedInterval,
    append_intervals_local,
    compute_bill_stats_from_intervals,
    extract_customer_info,
    extract_customer_info_from_files,
    load_intervals_local,
    parse_amr_file,
    parse_amr_files,
    remove_interval_rows,
    render_boxplot_png,
    summarize_available,
    summarize_available_by_account,
    update_account_metadata,
)


def _make_amr_html(start_date: dt.datetime, n_days: int, account_no: str = "", company_name: str = "") -> str:
    """สร้างไฟล์ HTML จำลองรูปแบบเดียวกับรายงาน "ข้อมูลกิโลวัตต์ชั่วโมงแบบช่วงเวลา" ของ PEA
    (ts + a1/b1/c1 ตาม P/OP/H) — ใช้ทดสอบ parse_amr_file โดยไม่ต้องมีไฟล์จริง

    ถ้าไม่ส่ง account_no/company_name จะได้ header ทั่วไปแบบเดิม — ส่งมาเพื่อจำลองหัวรายงานจริงของ
    PEA (ดู extract_customer_info)"""

    rows = []
    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):  # 96 x 15 นาที = 24 ชม.
            minutes = interval_i * 15
            t = d + dt.timedelta(minutes=minutes) + dt.timedelta(minutes=15)
            hour = (minutes // 60) % 24
            is_weekend = d.weekday() >= 5
            if is_weekend:
                col, kw = "c1", 100.0
            elif 9 <= hour < 22:
                col, kw = "a1", 400.0
            else:
                col, kw = "b1", 150.0
            row = {"ts": t.strftime("%d/%m/%Y %H.%M"), "a1": "", "a2": "", "b1": "", "b2": "", "c1": "", "c2": ""}
            row[col] = f"{kw:.2f}"
            rows.append(row)
        d += dt.timedelta(days=1)

    if account_no or company_name:
        header = (
            "<table width='800px' cellpadding='4' cellspacing='4'><tr>"
            "<td colspan='7' class='header'>รายงานข้อมูลกิโลวัตต์แบบช่วงเวลา</td></tr>"
            "<tr><td colspan='7' class='header'>[ระหว่างวันที่ : 01/01/2026 - 03/01/2026]</td></tr>"
            "<tr>"
            f"<td class='detail'>บัญชีผู้ใช้ไฟ : </td><td>{account_no}&nbsp;</td>"
            f"<td class='detail'>ชื่อผู้ใช้ไฟ : </td><td>{company_name}</td>"
            "<td></td><td></td><td></td>"
            "</tr>"
            "<tr>"
            "<td class='detail'>เครื่องวัด : </td><td>METER123&nbsp;</td>"
            "<td class='detail'>Tariff : </td><td>3.2&nbsp;</td>"
            "<td></td><td></td><td></td>"
            "</tr>"
            "<tr>"
            "<td class='detail'>CT Ratio : </td><td>1&nbsp;</td>"
            "<td class='detail'>VT Ratio : </td><td>1&nbsp;</td>"
            "<td></td><td></td><td></td>"
            "</tr>"
            "</table>"
        )
    else:
        header = "<table><tr><td>Header info</td></tr></table>"
    data_rows = "".join(
        "<tr>" + "".join(f"<td>{r[c]}</td>" for c in ["ts", "a1", "a2", "b1", "b2", "c1", "c2"]) + "</tr>"
        for r in rows
    )
    # ใช้ <td> (ไม่ใช่ <th>) สำหรับแถวหัวตารางในตารางข้อมูล — pandas.read_html จะไม่ดึงแถวนี้ไปตั้ง
    # เป็นชื่อคอลัมน์อัตโนมัติ (เหมือนไฟล์ PEA จริงที่ table[1] มีแถวข้อความหัวข้อปนอยู่เป็นแถวแรก
    # ของข้อมูลจริงๆ ไม่ใช่ <th> — ตรงกับที่ _read_one ต้อง .iloc[1:] ตัดแถวนั้นทิ้งเอง)
    data_table = (
        "<table><tr><td>ts</td><td>a1</td><td>a2</td><td>b1</td><td>b2</td><td>c1</td><td>c2</td></tr>"
        + data_rows
        + "</table>"
    )
    footer = "<table><tr><td>Footer info</td></tr></table>"
    return header + data_table + footer


@pytest.fixture()
def amr_report_file(tmp_path):
    html = _make_amr_html(dt.datetime(2026, 1, 1), n_days=5)  # 1 ม.ค. 2569 เป็นวันพฤหัส
    path = tmp_path / "report.xls"
    path.write_text(html, encoding="utf-8")
    return path


def _make_amr_xlsx(tmp_path, start_date: dt.datetime, n_days: int, account_no: str = "", company_name: str = "") -> Path:
    """สร้างไฟล์ .xlsx จริง (Excel 2007+ binary) จำลองรูปแบบไฟล์ AMR จริงแบบที่ 2 ที่ผู้ใช้เจอ
    (นอกจาก HTML-in-.xls แบบเดิม) — 3 ชีต (header/ข้อมูลจริง/ท้ายรายงาน) พร้อมแถวขยะ "0,1,2,...,6"
    บนสุดของทุกชีตที่เจอในไฟล์ตัวอย่างจริง (ดู _strip_leading_index_row ใน amr_boxplot.py)"""

    from openpyxl import Workbook

    wb = Workbook()
    junk_row = list(range(7))

    header_ws = wb.active
    header_ws.title = "Sheet1"
    header_ws.append(junk_row)
    header_ws.append(["รายงานข้อมูลกิโลวัตต์แบบช่วงเวลา"] * 7)
    header_ws.append(["[ระหว่างวันที่ : 01/01/2026 - 03/01/2026]"] * 7)
    header_ws.append(["บัญชีผู้ใช้ไฟ :", account_no, "ชื่อผู้ใช้ไฟ :", company_name, None, None, None])
    header_ws.append(["หมายเลขมิเตอร์ :", "METER123", "Tariff :", "TOU", None, None, None])
    header_ws.append(["CT Ratio :", "100:5 A.", "VT Ratio :", "115000:115 V.", None, None, None])

    data_ws = wb.create_sheet("Sheet2")
    data_ws.append(junk_row)
    data_ws.append([None, "RATE A", "RATE A", "RATE B", "RATE B", "RATE C", "RATE C"])
    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):
            minutes = interval_i * 15
            t = d + dt.timedelta(minutes=minutes) + dt.timedelta(minutes=15)
            hour = (minutes // 60) % 24
            row = [t.strftime("%d/%m/%Y %H.%M"), None, None, None, None, None, None]
            if d.weekday() >= 5:
                row[5] = row[6] = "100.000"
            elif 9 <= hour < 22:
                row[1] = row[2] = "400.000"
            else:
                row[3] = row[4] = "150.000"
            data_ws.append(row)
        d += dt.timedelta(days=1)

    footer_ws = wb.create_sheet("Sheet3")
    footer_ws.append(junk_row)
    footer_ws.append(["***หมายเหตุ***"] * 7)

    path = tmp_path / "report.xlsx"
    wb.save(path)
    return path


def _make_amr_interval_kwh_single_rate_xlsx(
    tmp_path, start_date: dt.datetime, n_days: int, account_no: str = "", company_name: str = ""
) -> Path:
    """สร้างไฟล์ .xlsx จริงจำลองรายงาน "ข้อมูลกิโลวัตต์ชั่วโมงแบบช่วงเวลา" ของ PEA (รูปแบบไฟล์ AMR จริง
    แบบที่ 9 ที่เจอ — 3 ชีต (header/ข้อมูลจริง/ท้ายรายงาน) เหมือน _make_amr_xlsx แต่ชีตข้อมูลจริงมีแค่
    5 คอลัมน์ (ts, RATE A, RATE B, RATE C, ผลรวม — ค่าเดียวต่อ rate ไม่ใช่คู่แบบ _make_amr_xlsx ที่มี
    7 คอลัมน์) และหน่วยข้อมูลเป็น kWh ต้องแปลง ×4 (เหมือน _make_amr_monthly_kwh_xlsx) มีแถวขยะ
    "0,1,2,3,4" บนสุดของทุกชีตเหมือนไฟล์ตัวอย่างจริง"""

    from openpyxl import Workbook

    wb = Workbook()
    junk_row = list(range(5))

    header_ws = wb.active
    header_ws.title = "Sheet1"
    header_ws.append(junk_row)
    header_ws.append(["รายงานข้อมูลกิโลวัตต์ชั่วโมงแบบช่วงเวลา"] * 5)
    header_ws.append(["[ระหว่างวันที่ : 01/01/2026 - 03/01/2026]"] * 5)
    header_ws.append(["บัญชีผู้ใช้ไฟ :", account_no, "ชื่อผู้ใช้ไฟ :", company_name, None])
    header_ws.append(["หมายเลขมิเตอร์ :", "METER123", "CT Ratio :", "100:5 A.", None])

    data_ws = wb.create_sheet("Sheet2")
    data_ws.append(junk_row)
    data_ws.append([None, "RATE A", "RATE B", "RATE C", "ผลรวม"])
    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):
            minutes = interval_i * 15
            t = d + dt.timedelta(minutes=minutes) + dt.timedelta(minutes=15)
            hour = (minutes // 60) % 24
            # kWh สะสมของ 15 นาที (÷4 ของ kW เฉลี่ยที่ตั้งใจให้ได้หลังแปลง ×4 กลับ)
            if d.weekday() >= 5:
                col, kwh = 2, 25.0  # -> kW เฉลี่ยหลังแปลง = 100.0
            elif 9 <= hour < 22:
                col, kwh = 0, 100.0  # -> 400.0
            else:
                col, kwh = 1, 37.5  # -> 150.0
            row_vals = [None, None, None]
            row_vals[col] = kwh
            data_ws.append([t.strftime("%d/%m/%Y %H.%M"), *row_vals, kwh])
        d += dt.timedelta(days=1)

    footer_ws = wb.create_sheet("Sheet3")
    footer_ws.append(junk_row)
    footer_ws.append(["***หมายเหตุ***"] * 5)
    footer_ws.append([f"พิมพ์โดย : {account_no}", None, None, None, "วันที่พิมพ์ : 01/01/2026 00:00"])

    path = tmp_path / "report_interval_kwh.xlsx"
    wb.save(path)
    return path


def _make_amr_monthly_kwh_xlsx(
    tmp_path, start_date: dt.datetime, n_days: int, account_no: str = "", company_name: str = ""
) -> Path:
    """สร้างไฟล์ .xlsx จริงจำลองรายงาน "กิโลวัตต์ชั่วโมงรายเดือน" ของ PEA (รูปแบบไฟล์ AMR จริงแบบที่ 3
    ที่เจอ — ชีตเดียว รวม header/ข้อมูล/ท้ายรายงานปนกันหมด หน่วยข้อมูลเป็น kWh ไม่ใช่ kW เฉลี่ยแบบ 2
    รูปแบบแรก ต้องแปลง ×4 ก่อนใช้ — ดู _read_one_monthly_kwh ใน amr_boxplot.py) ไม่มีแถวขยะ
    "0,1,2,...,6" แบบไฟล์รูปแบบที่ 2 (เจอเฉพาะไฟล์รูปแบบนั้น)"""

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "kWh0000000000000101202612000"

    ws.append([None, "การไฟฟ้าส่วนภูมิภาค", None, None, None])
    ws.append([None, "ที่อยู่การไฟฟ้าส่วนภูมิภาค", None, None, None])
    ws.append(["รายงานข้อมูลกิโลวัตต์ชั่วโมงรายเดือน", None, None, None, None])
    ws.append(["[ประจำเดือน : มกราคม 2569]", None, None, None, None])
    ws.append(["บัญชีผู้ใช้ไฟ :", f"{account_no}\xa0", "ชื่อผู้ใช้ไฟ :", company_name, None])
    ws.append(["หมายเลขมิเตอร์ :", "METER123\xa0", "Tariff :", "TOU", None])
    ws.append(["CT Ratio :", "100:5 A.\xa0", "VT Ratio :", "115000:115 V.", None])
    ws.append([None, None, None, None, None])
    ws.append([None, "RATE A", "RATE B", "RATE C", "ผลรวม"])

    totals = [0.0, 0.0, 0.0]
    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):
            minutes = interval_i * 15
            t = d + dt.timedelta(minutes=minutes) + dt.timedelta(minutes=15)
            hour = (minutes // 60) % 24
            # kWh สะสมของ 15 นาที (÷4 ของ kW เฉลี่ยที่ตั้งใจให้ได้หลังแปลง ×4 กลับ) — ใช้เลขกลมๆ
            # หาร 4 ลงตัวพอดี กันปัญหาเทียบ float เพี้ยนนิดหน่อยตอนเทส
            if d.weekday() >= 5:
                col, kwh = 2, 25.0  # -> kW เฉลี่ยหลังแปลง = 100.0
            elif 9 <= hour < 22:
                col, kwh = 0, 100.0  # -> 400.0
            else:
                col, kwh = 1, 37.5  # -> 150.0
            row_vals = [None, None, None]
            row_vals[col] = kwh
            totals[col] += kwh
            ws.append([f"\xa0{t.strftime('%d/%m/%Y %H.%M')}", *row_vals, kwh])
        d += dt.timedelta(days=1)

    ws.append(["ผลรวมทั้งหมด", totals[0], totals[1], totals[2], sum(totals)])
    ws.append(["***หมายเหตุ***", None, None, None, None])
    ws.append([None, None, None, None, None])
    ws.append([f"พิมพ์โดย : {account_no}", None, None, None, "วันที่พิมพ์ : 01/01/2026 00:00"])

    path = tmp_path / "report_monthly.xlsx"
    wb.save(path)
    return path


def _make_amr_monthly_kw_xlsx(
    tmp_path, start_date: dt.datetime, n_days: int, account_no: str = "", company_name: str = ""
) -> Path:
    """สร้างไฟล์ .xlsx จริงจำลองรายงาน "กิโลวัตต์รายเดือน" ของ PEA (รูปแบบไฟล์ AMR จริงแบบที่ 4 ที่
    เจอ — ชีตเดียวเหมือน _make_amr_monthly_kwh_xlsx แต่ชื่อรายงานไม่มีคำว่า "ชั่วโมง" และหน่วยข้อมูล
    เป็น kW เฉลี่ยตรงๆ อยู่แล้ว ไม่ต้องแปลงหน่วย — คอลัมน์เป็นคู่ a1/a2,b1/b2,c1/c2 แบบเดียวกับรายงาน
    "กิโลวัตต์แบบช่วงเวลา" แต่มีค่าจริงแค่คอลัมน์แรกของคู่ (ดู _read_one_monthly_kw)"""

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "kW0000000000000101202612000"

    ws.append(["รายงานข้อมูลกิโลวัตต์รายเดือน", None, None, None, None, None, None])
    ws.append(["[ประจำเดือน : มกราคม 2569]", None, None, None, None, None, None])
    ws.append(["บัญชีผู้ใช้ไฟ :", f"{account_no}\xa0", "ชื่อผู้ใช้ไฟ :", company_name, None, None, None])
    ws.append(["หมายเลขมิเตอร์ :", "METER123\xa0", "Tariff :", "TOU", None, None, None])
    ws.append(["CT Ratio :", "400:5 A.\xa0", "VT Ratio :", "22000:110 V.", None, None, None])
    ws.append([None, "RATE A", None, "RATE B", None, "RATE C", None])

    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):
            minutes = interval_i * 15
            t = d + dt.timedelta(minutes=minutes) + dt.timedelta(minutes=15)
            hour = (minutes // 60) % 24
            if d.weekday() >= 5:
                col, kw = 5, 100.0  # RATE C (H)
            elif 9 <= hour < 22:
                col, kw = 1, 400.0  # RATE A (P)
            else:
                col, kw = 3, 150.0  # RATE B (OP)
            row = [None] * 6
            row[col - 1] = kw
            ws.append([f"\xa0{t.strftime('%d/%m/%Y %H.%M')}", *row])
        d += dt.timedelta(days=1)

    ws.append(["กิโลวัตต์ต่ำสุด", 100.0, "01/01/2026 00.15", 150.0, "01/01/2026 00.15", 100.0, "01/01/2026 00.15"])
    ws.append(["กิโลวัตต์เฉลี่ย", 400.0, None, 150.0, None, 100.0, None])
    ws.append(["กิโลวัตต์สูงสุด", 400.0, "01/01/2026 09.15", 150.0, "01/01/2026 00.15", 100.0, "01/01/2026 00.15"])
    ws.append(["***หมายเหตุ***", None, None, None, None, None, None])
    ws.append([None, None, None, None, None, None, None])
    ws.append([f"พิมพ์โดย : {account_no}", None, None, None, None, None, "วันที่พิมพ์ : 01/01/2026 00:00"])

    path = tmp_path / "report_monthly_kw.xlsx"
    wb.save(path)
    return path


def _make_amr_custom_kw_xlsx(tmp_path, start_date: dt.datetime, n_days: int, account_no: str = "") -> Path:
    """สร้างไฟล์ .xlsx จริงจำลองรายงาน "Custom kW Report" ของ PEA (รูปแบบไฟล์ AMR จริงแบบที่ 5 ที่
    เจอ — ภาษาอังกฤษล้วน ชีตเดียว มีแค่คอลัมน์ "Time"/"kW" ไม่มีคอลัมน์ rate แยกเลย และไม่มีชื่อบริษัท
    ให้เลย — ต้องเดา rate จากวัน/เวลาเอง ดู _classify_tou_rate) เลขบัญชีในไฟล์จริงเป็น text ล้วนไม่มี
    ตัวอักษรอื่นปนเลยสักแถว (ต่างจากรูปแบบอื่นที่มี \\xa0 ต่อท้ายเสมอ) เพื่อจำลองบั๊กจริงที่เจอ (pandas
    เดา dtype คอลัมน์เป็น float ทำเลข 0 นำหน้าหาย) จึงตั้งใจไม่ใส่ \\xa0 ต่อท้ายในเทสนี้"""

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "kW0000000000000101202612000"

    ws.append([None, "Provincial Electricity Authority Advanced Metering Infrastructure (AMI) ", None])
    ws.append([None, "200 Ngam Wong Wan Road", None])
    ws.append([None, "Custom kW Report", None])
    ws.append([None, "Between 01 January 2026 - 31 January 2026", None])
    ws.append([None, "Contact Account : ", account_no])
    ws.append([None, "Meter No. : ", "6500651587"])
    ws.append([None, None, None])
    ws.append(["Time", "kW", None])

    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):
            minutes = interval_i * 15
            t = d + dt.timedelta(minutes=minutes) + dt.timedelta(minutes=15)
            ws.append([t.strftime("%d/%m/%Y %H.%M"), 123, None])
        d += dt.timedelta(days=1)

    path = tmp_path / "report_custom_kw.xlsx"
    wb.save(path)
    return path


def _make_amr_mea_csv(tmp_path, start_date: dt.datetime, n_days: int, mea_no: str = "140001877") -> Path:
    """สร้างไฟล์ CSV จำลองรูปแบบเดียวกับรายงาน AMR ของ กฟน. (MEA) จริง (รูปแบบไฟล์ AMR จริงแบบที่ 6
    ที่เจอ — คนละหน่วยงานกับ PEA ทั้งหมดที่รองรับอยู่ก่อนหน้านี้) คอลัมน์: MEA No., UI ID.,
    Rate Category, Measuring Compoenent, TOU/TOD, Date Time, Value — ปนข้อมูล kW (E-MAX-KW-IMP)
    กับ kVAR (E-MAX-KVAR-IMP) เข้าด้วยกันเหมือนไฟล์จริง (ต้องกรองเอาเฉพาะ kW) timestamp บอกเวลา
    "สิ้นสุด" ของช่วง 15 นาทีนั้น (เหมือนรายงานของ PEA ทุกแบบ) รูปแบบวันที่ M/D/YYYY H:MM (เดือนขึ้น
    ก่อนวัน ไม่เติมเลข 0 นำหน้า — ตามไฟล์ตัวอย่างจริง) TOU/TOD คำนวณจากกฎเดียวกับที่ไฟล์จริงใช้จริง
    (วันทำการ ช่วง 9:00-22:00 ของเวลา "เริ่มต้น" ของช่วง = ON นอกนั้น = OFF)"""

    lines = ["MEA No.,UI ID.,Rate Category,Measuring Compoenent,TOU/TOD,Date Time,Value"]
    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):
            end = d + dt.timedelta(minutes=(interval_i + 1) * 15)
            start = end - dt.timedelta(minutes=15)
            is_peak = start.weekday() < 5 and 9 <= start.hour < 22
            tou = "ON" if is_peak else "OFF"
            ts = f"{end.month}/{end.day}/{end.year} {end.hour}:{end.minute:02d}"
            lines.append(f"{mea_no},96131109,3.2,E-MAX-KW-IMP,{tou},{ts},123")
            lines.append(f"{mea_no},96131109,3.2,E-MAX-KVAR-IMP,{tou},{ts},5")
        d += dt.timedelta(days=1)

    path = tmp_path / "mea_report.csv"
    path.write_text("\n".join(lines), encoding="utf-8-sig")
    return path


def _make_amr_mea_wide_excel(tmp_path, start_date: dt.datetime, n_days: int, mea_no: str = "140002293") -> Path:
    """สร้างไฟล์ .xlsx จริงจำลองรายงาน AMR ของ กฟน. (MEA) รูปแบบ "ตารางแสดงผลข้อมูลแบบสรุปแนวขวาง"
    (รูปแบบไฟล์ AMR จริงแบบที่ 7 ที่เจอ — คนละรูปแบบไฟล์กับ MEA CSV แม้จะเป็นหน่วยงานเดียวกัน) มีหัว
    ตาราง 4 แถวก่อนถึงแถวข้อมูลจริง (ชื่อรายงาน/ว่าง/"Value"/หัวตารางจริง) 1 แถวมีค่า kW/kVAR/kWh/
    kVARh ของช่วง 15 นาทีเดียวกันครบในแถวเดียว (ไม่ต้องกรองแถวตามชนิดข้อมูลเหมือน MEA CSV) Date Time
    เป็นรูปแบบ ISO "YYYY-MM-DD HH:MM:SS" ตรงๆ (ต่างจาก MEA CSV ที่เป็น M/D/YYYY H:MM)"""

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(["ตารางแสดงผลข้อมูลแบบสรุปแนวขวาง"])
    ws.append([])
    ws.append([None, None, None, None, None, "Value", None, None, None])
    ws.append(["MEA No.", "UI ID.", "Rate Category", "Date Time", "TOU/TOD", "E-KVARH-TOTAL-IMP", "E-KWH-TOTAL-IMP", "E-MAX-KVAR-IMP", "E-MAX-KW-IMP"])

    d = start_date
    for _ in range(n_days):
        for interval_i in range(96):
            end = d + dt.timedelta(minutes=(interval_i + 1) * 15)
            start = end - dt.timedelta(minutes=15)
            is_peak = start.weekday() < 5 and 9 <= start.hour < 22
            tou = "ON" if is_peak else "OFF"
            ws.append([mea_no, "0095856820", 3.2, end.strftime("%Y-%m-%d %H:%M:%S"), tou, 0, 30.886, 0, 123.544])
        d += dt.timedelta(days=1)

    path = tmp_path / "mea_wide_report.xlsx"
    wb.save(path)
    return path


def _make_amr_power_logger_csv(tmp_path, samples: list, delimiter: str = "\t") -> Path:
    """สร้างไฟล์ .csv จำลองไฟล์ log จากเครื่องวัดไฟฟ้า (power logger — รุ่น DW-680) จริง (รูปแบบไฟล์
    AMR จริงแบบที่ 10 ที่เจอ — คนละประเภทไฟล์โดยสิ้นเชิงจากรายงาน AMR ของ PEA/MEA: preamble ตั้งค่า
    เครื่องช่วงบนจำนวนคอลัมน์ต่อแถวไม่เท่ากันเลย ตามด้วยหัวตาราง 2 แถว (กลุ่ม/ชื่อฟิลด์) แล้วข้อมูล
    วัดต่อเนื่องทุกไม่กี่วินาที ไม่ใช่ราย 15 นาที) samples คือ list ของ (date_str "YYYY/MM/DD" —
    ยืนยันจากไฟล์ตัวอย่างจริงตรงๆ ไม่ใช่รูปแบบที่ Excel แสดงผลใหม่ให้ตอนเปิดดู, time_str "H:M:S",
    psum_value) — ตั้งใจให้ทดสอบกำหนดเองได้ทั้งวันที่/เวลา/ค่า เพื่อเทียบผลลัพธ์
    การเฉลี่ยรวมเป็นช่วง 15 นาที (ดู _parse_power_logger_csv) delimiter เลือกได้ทั้ง tab (ตามที่วาง
    จาก Excel ปกติ) หรือ comma (ไฟล์ .csv มาตรฐาน) ให้ทดสอบว่า csv.Sniffer เดาได้ถูกทั้งคู่"""

    lines = [
        delimiter.join(["Model", "DW-680"]),
        delimiter.join(["Serial Number", "3524325002"]),
        "",
        delimiter.join(["User Name", "User01"]),
        delimiter.join(["Location", "EK1"]),
        "",
        delimiter.join(["Wire Mode", "3P4W_3CT"]),
        delimiter.join(["Frequency(Hz)", "50"]),
        delimiter.join(["Nominal Voltage(V)", "220"]),
        "",
        delimiter.join(["Date", "Time", "Phase Voltage(V)"]),
        delimiter.join(["", "", "UA", "UB", "UC", "PA", "PB", "PC", "PSum"]),
    ]
    for date_str, time_str, psum in samples:
        lines.append(delimiter.join([date_str, time_str, "233", "236", "235", "100", "100", "100", str(psum)]))

    path = tmp_path / "power_logger.csv"
    path.write_text("\n".join(lines), encoding="utf-8-sig")
    return path


def _make_amr_pea_monthly_combined_datetime_xls(tmp_path, rows: list) -> Path:
    """สร้างไฟล์ .xls ไบนารีเก่าจริง (BIFF/OLE2 — ใช้ xlwt เขียน ไม่ใช่ openpyxl ที่เขียนได้แค่ .xlsx)
    จำลองรายงาน "รายงานใช้ไฟ - รายเดือน" ของ PEA (รูปแบบไฟล์ AMR จริงแบบที่ 11 ที่เจอ — คนละไบนารี
    ฟอร์แมตโดยสิ้นเชิงจากทั้ง .xlsx จริงและ HTML-in-.xls แบบเดิม) คอลัมน์ "วันที่/เวลา" รวมวันที่+
    เวลาไว้คอลัมน์เดียว ปี พ.ศ. มีแค่ 2 คอลัมน์ kW (On-Peak/Off-Peak) ไม่มี Holiday แยก rows คือ
    list ของ (date_str "DD/MM/YYYY" ปี พ.ศ., time_str "HH:MM" หรือ "24:00", on_peak, off_peak)"""

    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet0")
    ws.write(0, 0, "รายงานใช้ไฟ - รายเดือน")
    ws.write(1, 0, "เครื่องวัดฯ: 73006874")
    ws.write(2, 0, "เงื่อนไข: เดือน มกราคม 2569")
    ws.write(4, 0, "วันที่/เวลา")
    ws.write(4, 1, "kW (On-Peak)")
    ws.write(4, 2, "kW (Off-Peak)")
    ws.write(4, 3, "kVAR (On-Peak)")
    ws.write(4, 4, "kVAR (Off-Peak)")
    for i, (date_str, time_str, on_peak, off_peak) in enumerate(rows):
        ws.write(5 + i, 0, f"{date_str} {time_str}")
        ws.write(5 + i, 1, on_peak)
        ws.write(5 + i, 2, off_peak)
        ws.write(5 + i, 3, 0)
        ws.write(5 + i, 4, 0)

    wb.add_sheet("Sheet1")  # ชีตว่างเปล่า เหมือนไฟล์ตัวอย่างจริง (Sheet1 ไม่มีข้อมูลเลย)

    path = tmp_path / "pea_monthly_combined.xls"
    wb.save(str(path))
    return path


def _make_amr_pea_monthly_combined_datetime_non_tou_xls(tmp_path, rows: list) -> Path:
    """เหมือน _make_amr_pea_monthly_combined_datetime_xls ทุกประการ แต่จำลองบัญชีที่ถือสัญญาอัตรา
    Non-TOU (ไม่มีการแยกช่วงเวลาเลย — เจอจากไฟล์ตัวอย่างจริงอีกบัญชีหนึ่ง) มีแค่คอลัมน์ "kW" เดี่ยวๆ
    ไม่มี On-Peak/Off-Peak แยก rows คือ list ของ (date_str "DD/MM/YYYY" ปี พ.ศ., time_str "HH:MM"
    หรือ "24:00", kw)"""

    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sheet0")
    ws.write(0, 0, "รายงานใช้ไฟ - รายเดือน")
    ws.write(1, 0, "เครื่องวัดฯ: 55051721")
    ws.write(2, 0, "เงื่อนไข: เดือน มกราคม 2569")
    ws.write(4, 0, "วันที่/เวลา")
    ws.write(4, 1, "kW")
    ws.write(4, 2, "kVAR")
    for i, (date_str, time_str, kw) in enumerate(rows):
        ws.write(5 + i, 0, f"{date_str} {time_str}")
        ws.write(5 + i, 1, kw)
        ws.write(5 + i, 2, 0)

    wb.add_sheet("Sheet1")

    path = tmp_path / "pea_monthly_combined_non_tou.xls"
    wb.save(str(path))
    return path


def _make_amr_load_profile_detail_xlsx(tmp_path, start_date: dt.datetime, n_days: int, kw: float = 500.0) -> Path:
    """สร้างไฟล์ .xlsx จริงจำลองรายงาน "Load Profile" (รูปแบบไฟล์ AMR จริงแบบที่ 8 ที่เจอ — ไฟล์
    "สรุปแล้ว" ไม่ใช่ export ดิบจาก PEA/MEA ตรงๆ ไม่มีเลขบัญชี/ชื่อบริษัทเลย) หัวตารางแถวเดียว (แถว 0)
    คอลัมน์ DateTime + kW รวม + แยกตาม TOD rate (On-Peak/Off-Peak/Partial-Peak) ที่รวมกันเท่ากับ
    คอลัมน์ kW เป๊ะเหมือนไฟล์จริง — สัญญาอัตรา TOD ใช้ 3 ช่วงราคาเดียวกันทุกวันในสัปดาห์ (ไม่แยกวันหยุด
    สุดสัปดาห์แบบ TOU) จึงตั้งทุกแถวเป็น On-Peak ล้วนให้ง่ายต่อการเทียบผลเทส (ระบบไม่ได้ใช้คอลัมน์นี้
    อยู่แล้ว — ดู _parse_load_profile_detail_xlsx)"""

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(
        [
            "olf", "Date", "Time", "DateTime", "Month", "Days", "TimeGroup", "kW",
            "RATE_A", "RATE_B", "RATE_C", "kW (On-Peak)", "kW (Off-Peak)", "kW (Partial-Peak)",
            "kVAR (On-Peak)", "kVAR (Off-Peak)", "kVAR (Partial-Peak)",
        ]
    )

    d = start_date
    olf = 1
    for _ in range(n_days):
        for interval_i in range(96):
            end = d + dt.timedelta(minutes=(interval_i + 1) * 15)
            ws.append(
                [
                    olf, end.strftime("%d/%m/%Y"), end.strftime("%H:%M"), end.strftime("%Y-%m-%d %H:%M:%S"),
                    end.strftime("%B"), end.strftime("%A"), end.strftime("%H:00"), kw,
                    kw, 0, 0, kw, 0, 0, 0, 0, 0,
                ]
            )
            olf += 1
        d += dt.timedelta(days=1)

    path = tmp_path / "load_profile_detail.xlsx"
    wb.save(path)
    return path


def test_parse_amr_file_extracts_all_intervals(amr_report_file):
    intervals = parse_amr_file(amr_report_file)
    assert len(intervals) == 5 * 96  # 5 วัน x 96 จุดต่อวัน ไม่มีจุดไหนถูกทิ้ง


def test_parse_amr_file_classifies_rate_correctly(amr_report_file):
    intervals = parse_amr_file(amr_report_file)
    rates = {i.rate for i in intervals}
    assert rates == {"P", "OP", "H"}


def test_parse_amr_file_converts_buddhist_era_year_to_gregorian(tmp_path):
    """ไฟล์ AMR จริงบางไฟล์ (โดยเฉพาะรายงาน AMI) ใช้ปี พ.ศ. (เช่น 2569) ในคอลัมน์วันที่แทน ค.ศ.
    (2026) — ต้องแปลงกลับเป็น ค.ศ. ให้ถูกต้อง ไม่งั้นวันที่ที่เก็บไว้จะผิดไป 543 ปี"""

    header = "<table><tr><td>Header info</td></tr></table>"
    data_table = (
        "<table><tr><td>ts</td><td>a1</td><td>a2</td><td>b1</td><td>b2</td><td>c1</td><td>c2</td></tr>"
        "<tr><td>01/01/2569 09.15</td><td>400.00</td><td>400.00</td><td></td><td></td><td></td><td></td></tr>"
        "</table>"
    )
    footer = "<table><tr><td>Footer info</td></tr></table>"
    path = tmp_path / "report_be_year.xls"
    path.write_text(header + data_table + footer, encoding="utf-8")

    intervals = parse_amr_file(path)
    assert len(intervals) == 1
    assert intervals[0].date == "2026-01-01"


def test_parse_amr_file_unreadable_returns_empty_list(tmp_path):
    bad_file = tmp_path / "not_amr.xls"
    bad_file.write_text("<html><body>not a real report</body></html>", encoding="utf-8")
    assert parse_amr_file(bad_file) == []


def test_parse_amr_files_dedupes_identical_files(amr_report_file):
    """อัปโหลดไฟล์เดิมซ้ำ 2 รอบ ต้องนับแค่ครั้งเดียว ไม่ใช่ x2 (เดิมเคยมีบั๊ก dedup ผิดระดับที่ทำให้
    เหลือแค่ 1 จุดต่อชั่วโมงแทนที่จะเป็น 4 จุด — เทสต์นี้ครอบคลุมทั้งสองแบบ)"""

    once = parse_amr_files([amr_report_file])
    twice = parse_amr_files([amr_report_file, amr_report_file])
    assert len(once) == len(twice)
    assert len(once) == 5 * 96


def test_append_and_load_intervals_local_round_trip(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    added = append_intervals_local("55101", intervals, storage)
    assert added == len(intervals)

    df = load_intervals_local(storage, "55101")
    assert len(df) == len(intervals)

    # ประเภทธุรกิจอื่นต้องไม่เห็นข้อมูลนี้
    other = load_intervals_local(storage, "86101")
    assert other.empty


def test_compute_bill_stats_from_intervals_matches_fixture(amr_report_file):
    """amr_report_file = 5 วันเริ่มพฤหัส (พฤ/ศุกร์/เสาร์/อาทิตย์/จันทร์) → 3 วันทำการ (P peak=400kW,
    OP peak=150kW) + 2 วันหยุด (H peak=100kW) — เช็คว่าคำนวณ peak/หน่วยไฟ/จำนวนวันตรงเป๊ะ"""

    stats = compute_bill_stats_from_intervals(parse_amr_file(amr_report_file))

    assert stats["P"] == {"peak": 400.0, "energy_kwh": 15600.0, "days": 3}
    assert stats["OP"] == {"peak": 150.0, "energy_kwh": 4950.0, "days": 3}
    assert stats["H"] == {"peak": 100.0, "energy_kwh": 4800.0, "days": 2}


def test_compute_bill_stats_from_intervals_empty_input_returns_empty_dict():
    assert compute_bill_stats_from_intervals([]) == {}


def test_load_intervals_local_missing_file_returns_empty(tmp_path):
    df = load_intervals_local(tmp_path / "nope.csv")
    assert df.empty


def test_load_intervals_local_filters_by_account_no(amr_report_file, tmp_path):
    """account_no=None (ค่าเริ่มต้น) ต้องรวมทุกบัญชี ส่วนระบุ account_no มาต้องกรองเหลือเฉพาะบัญชี
    นั้น — เคยเป็นบั๊กที่ /api/forecast-boxplot กรองแค่ business_type_code ทำให้กดดูกราฟของคนละบัญชี
    ใน TSIC เดียวกันแล้วได้กราฟเหมือนกันเป๊ะ"""

    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("55101", intervals, storage, account_no="A1")
    append_intervals_local("55101", intervals, storage, account_no="A2")

    all_df = load_intervals_local(storage, "55101")
    assert len(all_df) == len(intervals) * 2

    a1_df = load_intervals_local(storage, "55101", account_no="A1")
    assert len(a1_df) == len(intervals)
    assert (a1_df["account_no"] == "A1").all()

    a2_df = load_intervals_local(storage, "55101", account_no="A2")
    assert len(a2_df) == len(intervals)
    assert (a2_df["account_no"] == "A2").all()


def test_load_intervals_local_filters_unspecified_account_group(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("55101", intervals, storage)  # ไม่ระบุ account_no
    append_intervals_local("55101", intervals, storage, account_no="A1")

    unspecified_df = load_intervals_local(storage, "55101", account_no="")
    assert len(unspecified_df) == len(intervals)
    assert (unspecified_df["account_no"] == "").all()


def test_append_intervals_local_stores_account_and_company_name(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local(
        "55101", intervals, storage,
        account_no="0199000001", company_name="บริษัท ทดสอบ จำกัด", registration_no="0105544000157",
    )

    df = load_intervals_local(storage, "55101")
    assert (df["account_no"] == "0199000001").all()
    assert (df["company_name"] == "บริษัท ทดสอบ จำกัด").all()
    assert (df["registration_no"] == "0105544000157").all()


def test_append_intervals_local_defaults_account_fields_to_blank(amr_report_file, tmp_path):
    """ไม่ระบุ account_no/company_name เลย (เหมือนโหมดแนบไฟล์เองแบบเดิม) ต้องยังใช้งานได้ปกติ
    ไม่ error เว้นว่างไว้เฉยๆ"""

    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("55101", intervals, storage)

    df = load_intervals_local(storage, "55101")
    assert (df["account_no"] == "").all()
    assert (df["company_name"] == "").all()


def test_append_intervals_local_migrates_old_header_format(tmp_path):
    """ไฟล์เก่าก่อนเพิ่มคอลัมน์ account_no/company_name (header มีแค่ 5 คอลัมน์เดิม) ต้อง migrate
    อัตโนมัติตอน append แถวใหม่ — แถวเก่ายังอยู่ครบ (account_no/company_name เว้นว่างไว้เพราะย้อน
    กลับไปหาไม่ได้) แถวใหม่มีค่าที่ใส่มาถูกต้อง"""

    storage = tmp_path / "storage.csv"
    storage.write_text(
        "business_type_code,date,hour,rate,kw\n55101,2026-01-01,9,P,123.4\n", encoding="utf-8"
    )

    new_intervals = [ParsedInterval(date="2026-01-02", hour=10, minute=0, rate="P", kw=200.0)]
    append_intervals_local("55101", new_intervals, storage, account_no="0199000002", company_name="บริษัท ใหม่ จำกัด")

    df = load_intervals_local(storage, "55101")
    assert len(df) == 2
    old_row = df[df["date"] == "2026-01-01"].iloc[0]
    assert old_row["account_no"] == ""
    assert old_row["company_name"] == ""
    new_row = df[df["date"] == "2026-01-02"].iloc[0]
    assert new_row["account_no"] == "0199000002"
    assert new_row["company_name"] == "บริษัท ใหม่ จำกัด"


def test_summarize_available_by_account_groups_per_tsic_and_account(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)  # 5*96 = 480 intervals

    # บัญชีเดียวกัน อัปโหลด 2 รอบ (คนละไฟล์/คนละเดือน) ต้องรวมเป็นกลุ่มเดียวกัน
    append_intervals_local("55101", intervals[:100], storage, account_no="A1", company_name="บริษัท เอ", registration_no="0105544000157")
    append_intervals_local("55101", intervals[100:], storage, account_no="A1", company_name="บริษัท เอ")
    # อีกบัญชีของ TSIC เดียวกัน
    append_intervals_local("55101", intervals[:50], storage, account_no="A2", company_name="บริษัท บี")
    # ไม่ระบุบัญชีเลย (โหมดแนบไฟล์เองแบบเก่า) — ต้องรวมเป็น 1 กลุ่ม "ไม่ระบุบัญชี" แยกจาก A1/A2
    append_intervals_local("55101", intervals[:20], storage)

    rows = summarize_available_by_account(storage)
    by_account = {(r["business_type_code"], r["account_no"]): r for r in rows}

    assert by_account[("55101", "A1")]["intervals"] == 480
    assert by_account[("55101", "A1")]["company_name"] == "บริษัท เอ"
    assert by_account[("55101", "A1")]["registration_no"] == "0105544000157"  # เจอในรอบแรกพอ
    assert by_account[("55101", "A2")]["intervals"] == 50
    assert by_account[("55101", "A2")]["company_name"] == "บริษัท บี"
    assert by_account[("55101", "")]["intervals"] == 20
    assert by_account[("55101", "")]["company_name"] == ""

    # เรียงจากจุดข้อมูลเยอะไปน้อย
    assert [r["intervals"] for r in rows] == sorted((r["intervals"] for r in rows), reverse=True)


def test_summarize_available_by_account_missing_file_returns_empty_list(tmp_path):
    assert summarize_available_by_account(tmp_path / "nope.csv") == []


def test_remove_interval_rows_deletes_only_matching_tsic_and_account(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)  # 5*96 = 480 intervals

    append_intervals_local("55101", intervals[:100], storage, account_no="A1")
    append_intervals_local("55101", intervals[100:200], storage, account_no="A2")
    append_intervals_local("55101", intervals[200:250], storage)  # ไม่ระบุบัญชี
    append_intervals_local("86101", intervals[250:300], storage, account_no="A1")  # คนละ TSIC เลขบัญชีชนกัน

    removed = remove_interval_rows(storage, "55101", "A1")
    assert removed == 100

    rows = {(r["business_type_code"], r["account_no"]): r for r in summarize_available_by_account(storage)}
    assert ("55101", "A1") not in rows  # ลบไปแล้ว
    assert rows[("55101", "A2")]["intervals"] == 100  # บัญชีอื่นของ TSIC เดียวกันไม่โดนลบ
    assert rows[("55101", "")]["intervals"] == 50  # กลุ่มไม่ระบุบัญชีไม่โดนลบ
    assert rows[("86101", "A1")]["intervals"] == 50  # คนละ TSIC ไม่โดนลบแม้เลขบัญชีจะชนกัน


def test_remove_interval_rows_empty_account_no_deletes_unspecified_group_only(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)

    append_intervals_local("55101", intervals[:100], storage, account_no="A1")
    append_intervals_local("55101", intervals[100:150], storage)  # ไม่ระบุบัญชี

    removed = remove_interval_rows(storage, "55101", "")
    assert removed == 50

    rows = {(r["business_type_code"], r["account_no"]): r for r in summarize_available_by_account(storage)}
    assert ("55101", "") not in rows
    assert rows[("55101", "A1")]["intervals"] == 100


def test_remove_interval_rows_no_match_returns_zero(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    append_intervals_local("55101", parse_amr_file(amr_report_file), storage, account_no="A1")
    assert remove_interval_rows(storage, "86101", "A1") == 0


def test_remove_interval_rows_missing_file_returns_zero(tmp_path):
    assert remove_interval_rows(tmp_path / "nope.csv", "55101", "") == 0


def test_update_account_metadata_sets_company_name_and_registration_no(amr_report_file, tmp_path):
    """ใช้ตอนไฟล์ต้นทางไม่มีชื่อบริษัทให้เลย (เช่นรายงาน "Custom kW Report") — แก้ชื่อบริษัท/เลข
    ทะเบียนของบัญชีที่อัปโหลดไปแล้วได้โดยตรง ไม่ต้องลบแล้วอัปโหลดใหม่ทั้งก้อน"""

    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("86101", intervals[:100], storage, account_no="020027862234")
    append_intervals_local("86101", intervals[100:200], storage, account_no="A2")  # บัญชีอื่นไม่โดนแตะ

    updated = update_account_metadata(
        storage, "86101", "020027862234", company_name="บริษัท พริ้นซิเพิล เฮลท์แคร์ - มุกดาหาร จำกัด"
    )
    assert updated == 100

    rows = {(r["business_type_code"], r["account_no"]): r for r in summarize_available_by_account(storage)}
    assert rows[("86101", "020027862234")]["company_name"] == "บริษัท พริ้นซิเพิล เฮลท์แคร์ - มุกดาหาร จำกัด"
    assert rows[("86101", "A2")]["company_name"] == ""  # บัญชีอื่นไม่โดนแก้


def test_update_account_metadata_only_touches_given_fields(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("86101", intervals[:50], storage, account_no="A1", registration_no="0105544000157")

    update_account_metadata(storage, "86101", "A1", company_name="บริษัท ทดสอบ จำกัด")

    rows = {(r["business_type_code"], r["account_no"]): r for r in summarize_available_by_account(storage)}
    assert rows[("86101", "A1")]["company_name"] == "บริษัท ทดสอบ จำกัด"
    assert rows[("86101", "A1")]["registration_no"] == "0105544000157"  # ไม่ส่งมา ต้องไม่ถูกแก้


def test_update_account_metadata_no_match_returns_zero(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    append_intervals_local("86101", parse_amr_file(amr_report_file), storage, account_no="A1")
    assert update_account_metadata(storage, "99999", "A1", company_name="ไม่มีจริง") == 0


def test_update_account_metadata_missing_file_returns_zero(tmp_path):
    assert update_account_metadata(tmp_path / "nope.csv", "55101", "A1", company_name="x") == 0


def test_summarize_available_reports_per_business_type(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("55101", intervals, storage)

    summary = summarize_available(storage)
    assert summary["55101"]["intervals"] == len(intervals)
    assert summary["55101"]["days"] == 5


def test_render_boxplot_png_returns_valid_png(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("55101", intervals, storage)
    df = load_intervals_local(storage, "55101")

    png = render_boxplot_png(df, business_type_code="55101", business_type_name="โรงแรม")
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_boxplot_png_with_subtitle_returns_valid_png(amr_report_file, tmp_path):
    """subtitle (ใช้ตอนกราฟถูกกรองเหลือบัญชีเดียว) ต้องไม่ทำให้วาดกราฟพัง"""

    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(amr_report_file)
    append_intervals_local("55101", intervals, storage, account_no="A1")
    df = load_intervals_local(storage, "55101", account_no="A1")

    png = render_boxplot_png(df, business_type_code="55101", business_type_name="โรงแรม", subtitle="บัญชี A1 — บริษัท ทดสอบ จำกัด")
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_boxplot_png_stats_line_handles_missing_rates(tmp_path):
    """ถ้าข้อมูลมีแค่ rate เดียว (เช่น เฉพาะวันหยุด H) บรรทัดสรุป Peak/เฉลี่ยต้องไม่พังเพราะ
    peaks/means dict มีแค่ key เดียว (ดู stats_parts loop ใน render_boxplot_png)"""

    html = _make_amr_html(dt.datetime(2026, 1, 3), n_days=2)  # 3-4 ม.ค. 2569 = เสาร์-อาทิตย์ล้วน
    path = tmp_path / "report.xls"
    path.write_text(html, encoding="utf-8")

    storage = tmp_path / "storage.csv"
    intervals = parse_amr_file(path)
    assert all(iv.rate == "H" for iv in intervals)
    append_intervals_local("55101", intervals, storage)
    df = load_intervals_local(storage, "55101")

    png = render_boxplot_png(df, business_type_code="55101", business_type_name="โรงแรม")
    assert png[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_boxplot_png_empty_dataframe_raises(amr_report_file, tmp_path):
    storage = tmp_path / "storage.csv"
    df = load_intervals_local(storage, "55101")  # ไม่มีไฟล์เลย -> DataFrame ว่าง
    with pytest.raises(ValueError):
        render_boxplot_png(df, business_type_code="55101")


def test_extract_customer_info_reads_account_and_company_from_header(tmp_path):
    html = _make_amr_html(
        dt.datetime(2026, 1, 1), n_days=1, account_no="0199000001", company_name="บริษัท ทดสอบ จำกัด"
    )
    path = tmp_path / "report.xls"
    path.write_text(html, encoding="utf-8")

    account_no, company_name = extract_customer_info(path)
    assert account_no == "0199000001"
    assert company_name == "บริษัท ทดสอบ จำกัด"


def test_extract_customer_info_returns_empty_when_not_found(amr_report_file):
    # amr_report_file ใช้ header ทั่วไป ไม่มี "บัญชีผู้ใช้ไฟ"/"ชื่อผู้ใช้ไฟ" ให้อ่าน
    account_no, company_name = extract_customer_info(amr_report_file)
    assert account_no == ""
    assert company_name == ""


def test_extract_customer_info_missing_file_returns_empty(tmp_path):
    account_no, company_name = extract_customer_info(tmp_path / "nope.xls")
    assert account_no == ""
    assert company_name == ""


def test_extract_customer_info_from_files_uses_first_match(tmp_path):
    empty_html = _make_amr_html(dt.datetime(2026, 1, 1), n_days=1)
    matched_html = _make_amr_html(
        dt.datetime(2026, 1, 1), n_days=1, account_no="0199000002", company_name="บริษัท สอง จำกัด"
    )
    path1 = tmp_path / "report1.xls"
    path1.write_text(empty_html, encoding="utf-8")
    path2 = tmp_path / "report2.xls"
    path2.write_text(matched_html, encoding="utf-8")

    account_no, company_name = extract_customer_info_from_files([path1, path2])
    assert account_no == "0199000002"
    assert company_name == "บริษัท สอง จำกัด"


def test_extract_customer_info_from_files_all_empty_returns_empty(amr_report_file):
    account_no, company_name = extract_customer_info_from_files([amr_report_file])
    assert account_no == ""
    assert company_name == ""


def test_parse_amr_file_supports_real_xlsx_binary_format(tmp_path):
    """ไฟล์ .xlsx จริง (Excel 2007+ binary) รูปแบบที่ 2 นอกจาก HTML-in-.xls เดิม — เจอจากไฟล์
    ตัวอย่างจริงของผู้ใช้ที่อัปโหลดไม่ได้ก่อนหน้านี้"""

    path = _make_amr_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96
    assert {i.rate for i in intervals} == {"P", "OP", "H"}


def test_extract_customer_info_supports_real_xlsx_binary_format(tmp_path):
    path = _make_amr_xlsx(
        tmp_path, dt.datetime(2026, 1, 1), n_days=1, account_no="0199000003", company_name="บริษัท สาม จำกัด"
    )
    account_no, company_name = extract_customer_info(path)
    assert account_no == "0199000003"
    assert company_name == "บริษัท สาม จำกัด"


def test_parse_amr_file_supports_monthly_kwh_format_and_converts_to_kw(tmp_path):
    """ไฟล์ .xlsx จริงรูปแบบที่ 3 "กิโลวัตต์ชั่วโมงรายเดือน" — ชีตเดียว หน่วยเป็น kWh สะสมต่อช่วง 15
    นาที ต้องถูกแปลงเป็น kW เฉลี่ย (×4) ให้หน่วยตรงกับรายงานแบบอื่นที่ระบบรองรับอยู่ก่อนแล้ว"""

    path = _make_amr_monthly_kwh_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96
    assert {i.rate for i in intervals} == {"P", "OP", "H"}

    by_rate = {i.rate: i.kw for i in intervals}
    assert by_rate["P"] == pytest.approx(400.0)
    assert by_rate["OP"] == pytest.approx(150.0)
    assert by_rate["H"] == pytest.approx(100.0)


def test_extract_customer_info_supports_monthly_kwh_format(tmp_path):
    path = _make_amr_monthly_kwh_xlsx(
        tmp_path, dt.datetime(2026, 1, 1), n_days=1, account_no="0199000004", company_name="บริษัท สี่ จำกัด"
    )
    account_no, company_name = extract_customer_info(path)
    assert account_no == "0199000004"
    assert company_name == "บริษัท สี่ จำกัด"


def test_parse_amr_file_supports_interval_kwh_single_rate_format_and_converts_to_kw(tmp_path):
    """ไฟล์ .xlsx จริงรูปแบบที่ 9 "ข้อมูลกิโลวัตต์ชั่วโมงแบบช่วงเวลา" — 3 ชีตเหมือนรูปแบบที่ 2 เดิม
    แต่ชีตข้อมูลจริงมีแค่ 5 คอลัมน์ (ts + RATE A/B/C ค่าเดียว + ผลรวม ไม่ใช่คู่ a1/a2 แบบรูปแบบที่ 2)
    หน่วยเป็น kWh สะสมต่อช่วง 15 นาที ต้องถูกแปลงเป็น kW เฉลี่ย (×4)"""

    path = _make_amr_interval_kwh_single_rate_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96
    assert {i.rate for i in intervals} == {"P", "OP", "H"}

    by_rate = {i.rate: i.kw for i in intervals}
    assert by_rate["P"] == pytest.approx(400.0)
    assert by_rate["OP"] == pytest.approx(150.0)
    assert by_rate["H"] == pytest.approx(100.0)


def test_extract_customer_info_supports_interval_kwh_single_rate_format(tmp_path):
    path = _make_amr_interval_kwh_single_rate_xlsx(
        tmp_path, dt.datetime(2026, 1, 1), n_days=1, account_no="0199000009", company_name="บริษัท เก้า จำกัด"
    )
    account_no, company_name = extract_customer_info(path)
    assert account_no == "0199000009"
    assert company_name == "บริษัท เก้า จำกัด"


def test_parse_amr_file_supports_monthly_kw_format_without_unit_conversion(tmp_path):
    """ไฟล์ .xlsx จริงรูปแบบที่ 4 "กิโลวัตต์รายเดือน" (ไม่มีคำว่า "ชั่วโมง" ต่างจากรูปแบบที่ 3) —
    หน่วยเป็น kW เฉลี่ยตรงๆ อยู่แล้ว ต้องไม่ถูกแปลง ×4 ผิดเหมือนรูปแบบ kWh"""

    path = _make_amr_monthly_kw_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96
    assert {i.rate for i in intervals} == {"P", "OP", "H"}

    by_rate = {i.rate: i.kw for i in intervals}
    assert by_rate["P"] == pytest.approx(400.0)
    assert by_rate["OP"] == pytest.approx(150.0)
    assert by_rate["H"] == pytest.approx(100.0)


def test_extract_customer_info_supports_monthly_kw_format(tmp_path):
    path = _make_amr_monthly_kw_xlsx(
        tmp_path, dt.datetime(2026, 1, 1), n_days=1, account_no="0199000005", company_name="บริษัท ห้า จำกัด"
    )
    account_no, company_name = extract_customer_info(path)
    assert account_no == "0199000005"
    assert company_name == "บริษัท ห้า จำกัด"


def test_parse_amr_file_supports_custom_kw_format_and_classifies_by_tou(tmp_path):
    """ไฟล์ .xlsx จริงรูปแบบที่ 5 "Custom kW Report" (ภาษาอังกฤษล้วน) ไม่มีคอลัมน์ rate แยกให้เลย
    ต้องเดา P/OP/H จากวัน/เวลาเอง — 1 ม.ค. 2569 เป็นวันพฤหัส (วันทำการ) เป็นวันแรก 5 วัน ครอบคลุมทั้ง
    วันทำการและวันหยุดสุดสัปดาห์ (เสาร์-อาทิตย์)"""

    path = _make_amr_custom_kw_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96
    assert {i.rate for i in intervals} == {"P", "OP", "H"}

    by_key = {(i.date, i.hour): i.rate for i in intervals}
    assert by_key[("2026-01-01", 10)] == "P"  # พฤหัส 10:00 -> วันทำการ ช่วง Peak (9-22)
    assert by_key[("2026-01-01", 2)] == "OP"  # พฤหัส 02:00 -> วันทำการ นอกช่วง Peak
    assert by_key[("2026-01-03", 10)] == "H"  # เสาร์ -> วันหยุด ไม่ว่าจะกี่โมง
    assert by_key[("2026-01-04", 10)] == "H"  # อาทิตย์ -> วันหยุด


def test_extract_customer_info_supports_custom_kw_format_without_leading_zero_loss(tmp_path):
    """เลขบัญชีในไฟล์รูปแบบนี้เป็น text ล้วนไม่มีอักขระอื่นปนเลยสักแถว (ต่างจากรูปแบบอื่นที่มี \\xa0
    ต่อท้ายเสมอ) — เคยเป็นบั๊กจริง: pandas.read_excel เดา dtype คอลัมน์เป็น float เพราะค่าดูเหมือน
    ตัวเลขล้วน ทำให้เลข 0 นำหน้าหาย (เช่น "020027862234" กลายเป็น "20027862234.0")"""

    path = _make_amr_custom_kw_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=1, account_no="020027862234")
    account_no, company_name = extract_customer_info(path)
    assert account_no == "020027862234"
    assert company_name == ""


def test_parse_amr_file_supports_mea_csv_format_and_filters_kw_only(tmp_path):
    """ไฟล์ CSV ของ กฟน. (MEA) ปนข้อมูล kW กับ kVAR เข้าด้วยกัน — ต้องอ่านเฉพาะแถว kW
    (E-MAX-KW-IMP) เท่านั้น ไม่นับแถว kVAR (E-MAX-KVAR-IMP) เลย"""

    path = _make_amr_mea_csv(tmp_path, dt.datetime(2026, 4, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96  # ไม่ใช่ 5*96*2 (ไม่รวม kVAR)
    assert all(i.kw == 123.0 for i in intervals)


def test_parse_amr_file_classifies_mea_rate_from_tou_and_weekday(tmp_path):
    """1 เม.ย. 2026 เป็นวันพุธ (วันทำการ) — ทดสอบว่าจับคู่คอลัมน์ TOU/TOD (ON/OFF) ร่วมกับวันใน
    สัปดาห์เป็น P/OP/H ถูกต้อง ครอบคลุมทั้งวันทำการและวันหยุดสุดสัปดาห์ (เสาร์-อาทิตย์)"""

    path = _make_amr_mea_csv(tmp_path, dt.datetime(2026, 4, 1), n_days=6)
    intervals = parse_amr_file(path)
    assert {i.rate for i in intervals} == {"P", "OP", "H"}

    by_key = {(i.date, i.hour): i.rate for i in intervals}
    assert by_key[("2026-04-01", 10)] == "P"  # พุธ 10:00 -> วันทำการ ช่วง Peak (9-22)
    assert by_key[("2026-04-01", 2)] == "OP"  # พุธ 02:00 -> วันทำการ นอกช่วง Peak
    assert by_key[("2026-04-04", 10)] == "H"  # เสาร์ -> วันหยุด ไม่ว่าจะกี่โมง
    assert by_key[("2026-04-05", 10)] == "H"  # อาทิตย์ -> วันหยุด


def test_extract_customer_info_supports_mea_csv_reads_mea_no_no_company_name(tmp_path):
    """ไฟล์รูปแบบนี้มีแค่เลขบัญชี (คอลัมน์ MEA No.) ไม่มีชื่อบริษัทให้เลย — company_name ต้องว่างเสมอ"""

    path = _make_amr_mea_csv(tmp_path, dt.datetime(2026, 4, 1), n_days=1, mea_no="140001877")
    account_no, company_name = extract_customer_info(path)
    assert account_no == "140001877"
    assert company_name == ""


def test_parse_amr_file_supports_mea_wide_excel_format(tmp_path):
    """ไฟล์ .xlsx จริงรูปแบบ "ตารางแสดงผลข้อมูลแบบสรุปแนวขวาง" ของ กฟน. (MEA) — คนละรูปแบบไฟล์กับ
    MEA CSV แม้จะเป็นหน่วยงานเดียวกัน ต้องอ่านคอลัมน์ "E-MAX-KW-IMP" ได้ถูกต้อง ไม่ใช่คอลัมน์พลังงาน
    สะสม (kWh/kVARh) หรือ kVAR"""

    path = _make_amr_mea_wide_excel(tmp_path, dt.datetime(2026, 4, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96
    assert all(i.kw == 123.544 for i in intervals)


def test_parse_amr_file_classifies_mea_wide_excel_rate_from_tou_and_weekday(tmp_path):
    """1 เม.ย. 2026 เป็นวันพุธ (วันทำการ) — เหมือนเทส MEA CSV ทุกประการ แต่ใช้ไฟล์รูปแบบแนวขวางแทน"""

    path = _make_amr_mea_wide_excel(tmp_path, dt.datetime(2026, 4, 1), n_days=6)
    intervals = parse_amr_file(path)
    assert {i.rate for i in intervals} == {"P", "OP", "H"}

    by_key = {(i.date, i.hour): i.rate for i in intervals}
    assert by_key[("2026-04-01", 10)] == "P"  # พุธ 10:00 -> วันทำการ ช่วง Peak (9-22)
    assert by_key[("2026-04-01", 2)] == "OP"  # พุธ 02:00 -> วันทำการ นอกช่วง Peak
    assert by_key[("2026-04-04", 10)] == "H"  # เสาร์ -> วันหยุด ไม่ว่าจะกี่โมง
    assert by_key[("2026-04-05", 10)] == "H"  # อาทิตย์ -> วันหยุด


def test_extract_customer_info_supports_mea_wide_excel_reads_mea_no_no_company_name(tmp_path):
    """ไฟล์รูปแบบนี้ไม่มีชื่อบริษัทให้เลยเหมือน MEA CSV — company_name ต้องว่างเสมอ"""

    path = _make_amr_mea_wide_excel(tmp_path, dt.datetime(2026, 4, 1), n_days=1, mea_no="140002293")
    account_no, company_name = extract_customer_info(path)
    assert account_no == "140002293"
    assert company_name == ""


def test_parse_amr_file_supports_load_profile_detail_format(tmp_path):
    """ไฟล์ .xlsx "Load Profile" (ไม่ใช่ export ดิบจาก PEA/MEA — ไม่มีเลขบัญชี/ชื่อบริษัทเลย) ต้อง
    อ่านคอลัมน์ "kW" ตรงๆ ได้ถูกต้อง ไม่ต้องรวมเองจากคอลัมน์ TOD rate แยก"""

    path = _make_amr_load_profile_detail_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=5, kw=500.0)
    intervals = parse_amr_file(path)
    assert len(intervals) == 5 * 96
    assert all(i.kw == 500.0 for i in intervals)


def test_parse_amr_file_classifies_load_profile_detail_rate_from_date_time(tmp_path):
    """ไฟล์นี้ไม่มีคอลัมน์ rate P/OP/H ให้เลย (คอลัมน์ TOD rate ของไฟล์ไม่ตรงกับแนวคิด P/OP/H ของ
    ระบบนี้ — ดู _parse_load_profile_detail_xlsx) ต้องเดาจากวัน/เวลาเองเหมือนรายงาน "Custom kW
    Report" ของ PEA — 1 ม.ค. 2569 เป็นวันพฤหัส (วันทำการ) ครอบคลุมทั้งวันทำการและวันหยุดสุดสัปดาห์"""

    path = _make_amr_load_profile_detail_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=5)
    intervals = parse_amr_file(path)
    assert {i.rate for i in intervals} == {"P", "OP", "H"}

    by_key = {(i.date, i.hour): i.rate for i in intervals}
    assert by_key[("2026-01-01", 10)] == "P"  # พฤหัส 10:00 -> วันทำการ ช่วง Peak (9-22)
    assert by_key[("2026-01-01", 2)] == "OP"  # พฤหัส 02:00 -> วันทำการ นอกช่วง Peak
    assert by_key[("2026-01-03", 10)] == "H"  # เสาร์ -> วันหยุด ไม่ว่าจะกี่โมง
    assert by_key[("2026-01-04", 10)] == "H"  # อาทิตย์ -> วันหยุด


def test_extract_customer_info_returns_empty_for_load_profile_detail_format(tmp_path):
    """ไฟล์รูปแบบนี้ไม่มีเลขบัญชี/ชื่อบริษัทให้เลย — ต้องคืนค่าว่างทั้งคู่เสมอ ไม่ error"""

    path = _make_amr_load_profile_detail_xlsx(tmp_path, dt.datetime(2026, 1, 1), n_days=1)
    account_no, company_name = extract_customer_info(path)
    assert account_no == ""
    assert company_name == ""


def test_parse_amr_file_averages_power_logger_samples_into_15min_buckets(tmp_path):
    """ไฟล์ log จากเครื่องวัดไฟฟ้าวัดต่อเนื่องทุกไม่กี่วินาที (ไม่ใช่ราย 15 นาทีแบบรายงาน AMR) —
    ต้องเฉลี่ยรวมเป็นช่วง 15 นาทีเอง ไม่ใช่เก็บทุกจุดดิบตรงๆ (ไม่งั้นถ่วงน้ำหนักสถิติ Boxplot ผิดเพี้ยน)
    2026-01-07 เป็นวันพุธ (วันทำการ) — 3 ตัวอย่างในช่วง Peak (10:00-10:00:20) เฉลี่ยรวมเป็น bucket
    เดียว (10:00) และ 2 ตัวอย่างในช่วง Off-Peak (02:00) เฉลี่ยรวมเป็นอีก bucket (02:00)"""

    samples = [
        ("2026/01/07", "10:00:00", 100.0),
        ("2026/01/07", "10:00:10", 200.0),
        ("2026/01/07", "10:00:20", 300.0),
        ("2026/01/07", "2:00:00", 50.0),
        ("2026/01/07", "2:00:10", 150.0),
    ]
    path = _make_amr_power_logger_csv(tmp_path, samples)
    intervals = parse_amr_file(path)
    assert len(intervals) == 2  # 5 จุดดิบ รวมเหลือ 2 bucket ไม่ใช่ 5 interval แยกกัน

    by_key = {(i.date, i.hour, i.minute): (i.kw, i.rate) for i in intervals}
    assert by_key[("2026-01-07", 10, 0)] == (pytest.approx(200.0), "P")  # เฉลี่ย (100+200+300)/3
    assert by_key[("2026-01-07", 2, 0)] == (pytest.approx(100.0), "OP")  # เฉลี่ย (50+150)/2


def test_parse_amr_file_classifies_power_logger_holiday_from_date(tmp_path):
    """2026-01-10 เป็นวันเสาร์ — ต้องจัดเป็นวันหยุด (H) ไม่ว่าจะกี่โมง แม้จะอยู่ในช่วงเวลาที่ปกติ
    เป็น Peak (9-22) ของวันทำการก็ตาม"""

    samples = [("2026/01/10", "10:00:00", 500.0)]
    path = _make_amr_power_logger_csv(tmp_path, samples)
    intervals = parse_amr_file(path)
    assert len(intervals) == 1
    assert intervals[0].rate == "H"
    assert intervals[0].date == "2026-01-10"


def test_parse_amr_file_supports_power_logger_csv_with_comma_delimiter(tmp_path):
    """ไฟล์ .csv บางไฟล์ใช้ comma คั่นจริงๆ (ไม่ใช่ tab ที่วางจาก Excel) — csv.Sniffer ต้องเดาถูก
    ทั้งคู่"""

    samples = [("2026/01/07", "10:00:00", 400.0)]
    path = _make_amr_power_logger_csv(tmp_path, samples, delimiter=",")
    intervals = parse_amr_file(path)
    assert len(intervals) == 1
    assert intervals[0].kw == pytest.approx(400.0)
    assert intervals[0].rate == "P"


def test_extract_customer_info_returns_empty_for_power_logger_format(tmp_path):
    """ไฟล์รูปแบบนี้ไม่มีเลขบัญชี/ชื่อบริษัทที่เชื่อถือได้เลย — ต้องคืนค่าว่างทั้งคู่เสมอ ไม่ error"""

    path = _make_amr_power_logger_csv(tmp_path, [("2026/01/07", "10:00:00", 100.0)])
    account_no, company_name = extract_customer_info(path)
    assert account_no == ""
    assert company_name == ""


def test_parse_amr_file_supports_pea_monthly_combined_datetime_xls(tmp_path):
    """ไฟล์ .xls ไบนารีเก่าจริง (BIFF/OLE2) รายงาน "รายงานใช้ไฟ - รายเดือน" ของ PEA — คอลัมน์
    "วันที่/เวลา" รวมกัน ปี พ.ศ. 2569 ต้องแปลงเป็น ค.ศ. 2026 ถูกต้อง 7 ม.ค. 2569(พ.ศ.) เป็นวันพุธ
    (วันทำการ) — ครอบคลุมทั้งช่วง On-Peak/Off-Peak ของวันทำการ"""

    rows = [
        ("07/01/2569", "10:15", 200.0, 0),  # ts บอกจุดสิ้นสุดของช่วง -> เริ่ม 10:00 วันทำการ ช่วง Peak
        ("07/01/2569", "02:15", 0, 80.0),  # เริ่ม 02:00 วันทำการ นอกช่วง Peak
    ]
    path = _make_amr_pea_monthly_combined_datetime_xls(tmp_path, rows)
    intervals = parse_amr_file(path)
    assert len(intervals) == 2
    assert {i.rate for i in intervals} == {"P", "OP"}

    by_key = {(i.date, i.hour): (i.rate, i.kw) for i in intervals}
    assert by_key[("2026-01-07", 10)] == ("P", pytest.approx(200.0))
    assert by_key[("2026-01-07", 2)] == ("OP", pytest.approx(80.0))


def test_parse_amr_file_classifies_pea_monthly_combined_datetime_holiday_from_weekday(tmp_path):
    """ไม่มีคอลัมน์ Holiday แยกให้เลย (มีแค่ On-Peak/Off-Peak) — ต้องเดา H จากวันในสัปดาห์เอง
    10 ม.ค. 2569(พ.ศ.) เป็นวันเสาร์ — ต้องจัดเป็นวันหยุด (H) แม้ On-Peak จะเป็น 0 (ค่าทั้งหมดอยู่ใน
    Off-Peak เพราะวันหยุดไม่มีช่วง Peak เลย)"""

    rows = [("10/01/2569", "10:00", 0, 150.0)]
    path = _make_amr_pea_monthly_combined_datetime_xls(tmp_path, rows)
    intervals = parse_amr_file(path)
    assert len(intervals) == 1
    assert intervals[0].rate == "H"
    assert intervals[0].date == "2026-01-10"
    assert intervals[0].kw == pytest.approx(150.0)


def test_parse_amr_file_handles_pea_monthly_combined_datetime_midnight_2400(tmp_path):
    """เวลา "24:00" หมายถึงเที่ยงคืนของวันถัดไป (pd.Timestamp ไม่รับ hour=24 ตรงๆ) — ต้องแปลงวันที่
    ให้ถูกต้อง ไม่ error/ตกหล่น"""

    rows = [("31/01/2569", "24:00", 0, 90.0)]
    path = _make_amr_pea_monthly_combined_datetime_xls(tmp_path, rows)
    intervals = parse_amr_file(path)
    assert len(intervals) == 1
    # timestamp "31/01/2569 24:00" = เที่ยงคืนของ 1 ก.พ. 2569(พ.ศ.) = 2026-02-01 00:00 (ค.ศ.)
    # ลบ 15 นาที (ts บอกจุดสิ้นสุดของช่วง) -> 2026-01-31 23:45 คือจุดเริ่มต้นของช่วงสุดท้าย
    assert intervals[0].date == "2026-01-31"
    assert intervals[0].hour == 23
    assert intervals[0].minute == 45


def test_extract_customer_info_returns_empty_for_pea_monthly_combined_datetime_xls(tmp_path):
    """ไฟล์รูปแบบนี้มีแค่หมายเลขมิเตอร์ (เครื่องวัดฯ) ไม่มีเลขบัญชีผู้ใช้ไฟ/ชื่อบริษัทให้เลย — ต้อง
    คืนค่าว่างทั้งคู่เสมอ ไม่ดึงหมายเลขมิเตอร์มาใช้แทนเลขบัญชี"""

    path = _make_amr_pea_monthly_combined_datetime_xls(tmp_path, [("07/01/2569", "10:00", 200.0, 0)])
    account_no, company_name = extract_customer_info(path)
    assert account_no == ""
    assert company_name == ""


def test_parse_amr_file_supports_pea_monthly_combined_datetime_non_tou(tmp_path):
    """บัญชีที่ถือสัญญาอัตรา Non-TOU ไม่มีคอลัมน์ On-Peak/Off-Peak แยกเลย มีแค่ "kW" เดี่ยวๆ — ต้อง
    เดา P/OP/H จากวัน/เวลาเองทั้งหมด (เหมือน "Custom kW Report" ของ PEA) 7 ม.ค. 2569(พ.ศ.) เป็น
    วันพุธ (วันทำการ) — ครอบคลุมทั้งช่วงที่ควรเป็น Peak/Off-Peak ตามเวลา แม้ไฟล์จะไม่บอกมาตรงๆ ก็ตาม"""

    rows = [
        ("07/01/2569", "10:15", 200.0),  # เริ่ม 10:00 วันทำการ ช่วง Peak (9-22)
        ("07/01/2569", "02:15", 80.0),  # เริ่ม 02:00 วันทำการ นอกช่วง Peak
    ]
    path = _make_amr_pea_monthly_combined_datetime_non_tou_xls(tmp_path, rows)
    intervals = parse_amr_file(path)
    assert len(intervals) == 2
    assert {i.rate for i in intervals} == {"P", "OP"}

    by_key = {(i.date, i.hour): (i.rate, i.kw) for i in intervals}
    assert by_key[("2026-01-07", 10)] == ("P", pytest.approx(200.0))
    assert by_key[("2026-01-07", 2)] == ("OP", pytest.approx(80.0))


def test_parse_amr_file_classifies_pea_monthly_combined_datetime_non_tou_holiday(tmp_path):
    """10 ม.ค. 2569(พ.ศ.) เป็นวันเสาร์ — ต้องจัดเป็นวันหยุด (H) แม้จะอยู่ในช่วงเวลาที่ปกติเป็น Peak
    ของวันทำการก็ตาม เพราะไฟล์ไม่มีข้อมูล rate ให้เลย ต้องอาศัยวันในสัปดาห์ล้วนๆ"""

    path = _make_amr_pea_monthly_combined_datetime_non_tou_xls(tmp_path, [("10/01/2569", "10:15", 300.0)])
    intervals = parse_amr_file(path)
    assert len(intervals) == 1
    assert intervals[0].rate == "H"
    assert intervals[0].date == "2026-01-10"


def test_extract_customer_info_returns_empty_for_pea_monthly_combined_datetime_non_tou(tmp_path):
    path = _make_amr_pea_monthly_combined_datetime_non_tou_xls(tmp_path, [("07/01/2569", "10:15", 200.0)])
    account_no, company_name = extract_customer_info(path)
    assert account_no == ""
    assert company_name == ""
