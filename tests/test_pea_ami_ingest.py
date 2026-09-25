import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from amr_mapping.pea_ami_ingest import parse_ami_interval_report, parse_ami_report_header
from amr_mapping.pea_ingest import IntervalReading, is_ami_xlsx, parse_interval_report, parse_report_header


# โครงสร้างจริงจากไฟล์ export ของระบบ "AMI" ของ PEA (ยืนยันจากไฟล์จริงที่ผู้ใช้ส่งมา) — ใช้
# ข้อมูลสมมติแทนของจริงทั้งหมด ปีในคอลัมน์เวลาเป็น พ.ศ. ตามไฟล์จริง (2568 = ค.ศ. 2025)
def _write_synthetic_ami_workbook(path, account_no="0199000000", company_name="บริษัท ทดสอบ จำกัด"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["B3"] = "การไฟฟ้าส่วนภูมิภาค โครงการติดตั้งระบบมิเตอร์อัจฉริยะ (AMI)"
    ws["B6"] = "รายงานข้อมูลกิโลวัตต์รายเดือน"
    ws["B7"] = "ประจำเดือน : กันยายน 2568"
    ws["B9"] = "บัญชีผู้ใช้ไฟฟ้า : "
    ws["C9"] = account_no
    ws["E9"] = "ชื่อผู้ใช้ไฟฟ้า : "
    ws["F9"] = company_name
    ws["B10"] = "หมายเลขมิเตอร์ : "
    ws["C10"] = "1234567"
    ws["E10"] = "Tariff : "
    ws["F10"] = "TOU"
    ws["B11"] = "CT Ratio : "
    ws["C11"] = "150:5 A."
    ws["E11"] = "VT Ratio : "
    ws["F11"] = "115000:115 V."

    ws["A28"] = "เวลา"
    ws["B28"] = "Rate A"
    ws["D28"] = "Rate B"
    ws["F28"] = "Rate C"
    ws["A29"] = "01/09/2568 00.15"
    ws["D29"] = "20.000"
    ws["A30"] = "01/09/2568 09.30"
    ws["B30"] = "30.000"
    ws["A31"] = "01/09/2568 22.15"
    ws["F31"] = "10.000"
    ws["A32"] = "กิโลวัตต์สูงสุด"
    ws["B32"] = "24/09/2568 18.00"
    ws["C32"] = "292.000"
    ws["A34"] = "พิมพ์โดย : ทดสอบ"

    wb.save(path)


def test_is_ami_xlsx_true_for_real_excel_file(tmp_path):
    path = tmp_path / "sample.xlsx"
    _write_synthetic_ami_workbook(path)
    assert is_ami_xlsx(path) is True


def test_is_ami_xlsx_false_for_html_disguised_as_xls(tmp_path):
    path = tmp_path / "sample.xls"
    path.write_text("<html><body>ไม่ใช่ Excel จริง</body></html>", encoding="utf-8")
    assert is_ami_xlsx(path) is False


def test_parse_ami_report_header_extracts_and_aliases_fields(tmp_path):
    path = tmp_path / "ami_report.xlsx"
    _write_synthetic_ami_workbook(path, account_no="0199000000", company_name="บริษัท ทดสอบ จำกัด")

    info = parse_ami_report_header(path)

    # ป้ายในไฟล์ AMI มีคำว่า "ไฟฟ้า" ต่อท้าย (บัญชีผู้ใช้ไฟฟ้า/ชื่อผู้ใช้ไฟฟ้า) ต้องถูกแปลงให้
    # ตรงกับ key มาตรฐานเดียวกับไฟล์ AMRWEB (บัญชีผู้ใช้ไฟ/ชื่อผู้ใช้ไฟ) ไม่งั้นโค้ดฝั่งเรียกใช้
    # (amr_import.py) จะหาไม่เจอ
    assert info["บัญชีผู้ใช้ไฟ"] == "0199000000"
    assert info["ชื่อผู้ใช้ไฟ"] == "บริษัท ทดสอบ จำกัด"
    assert info["หมายเลขมิเตอร์"] == "1234567"
    assert info["Tariff"] == "TOU"
    assert info["CT Ratio"] == "150:5 A."
    assert info["VT Ratio"] == "115000:115 V."


def test_parse_ami_interval_report_reads_rate_columns_and_converts_be_year(tmp_path):
    path = tmp_path / "ami_report.xlsx"
    _write_synthetic_ami_workbook(path)

    readings = parse_ami_interval_report(path)

    # ปี พ.ศ. 2568 ในไฟล์ ต้องถูกแปลงเป็น ค.ศ. 2025 ก่อนคืนค่า (ต่างจากไฟล์ AMRWEB ที่ใช้ ค.ศ.
    # อยู่แล้ว ไม่งั้น pea_ingest._parse_interval_timestamp จะคำนวณวันในสัปดาห์ผิด)
    assert IntervalReading(timestamp="01/09/2025 00.15", period="OP", kwh=20.0) in readings
    assert IntervalReading(timestamp="01/09/2025 09.30", period="P", kwh=30.0) in readings
    assert IntervalReading(timestamp="01/09/2025 22.15", period="H", kwh=10.0) in readings


def test_parse_ami_interval_report_skips_summary_and_footer_rows(tmp_path):
    path = tmp_path / "ami_report.xlsx"
    _write_synthetic_ami_workbook(path)

    readings = parse_ami_interval_report(path)

    # แถว "กิโลวัตต์สูงสุด" และ "พิมพ์โดย" ท้ายตาราง ต้องไม่ถูกอ่านเป็น reading (timestamp
    # ของแถวเหล่านี้ไม่ตรงรูปแบบวันที่ DD/MM/YYYY)
    assert len(readings) == 3
    assert all(r.timestamp.startswith("01/09/2025") for r in readings)


def test_pea_ingest_parse_report_header_dispatches_to_ami_parser_for_xlsx(tmp_path):
    """pea_ingest.parse_report_header (ที่ amr_import.py เรียกใช้อยู่แล้ว) ต้องตรวจจับไฟล์
    .xlsx แท้อัตโนมัติแล้วส่งต่อให้ parse_ami_report_header โดยผู้เรียกไม่ต้องรู้ว่าไฟล์เป็น
    รูปแบบไหน"""

    path = tmp_path / "ami_report.xlsx"
    _write_synthetic_ami_workbook(path, account_no="0288888888", company_name="บริษัท เอมี จำกัด")

    info = parse_report_header(path)

    assert info["บัญชีผู้ใช้ไฟ"] == "0288888888"
    assert info["ชื่อผู้ใช้ไฟ"] == "บริษัท เอมี จำกัด"


def test_pea_ingest_parse_interval_report_dispatches_to_ami_parser_for_xlsx(tmp_path):
    path = tmp_path / "ami_report.xlsx"
    _write_synthetic_ami_workbook(path)

    readings = parse_interval_report(path)

    assert IntervalReading(timestamp="01/09/2025 09.30", period="P", kwh=30.0) in readings


# โครงสร้างจริงอีกแบบหนึ่ง (ยืนยันจากไฟล์จริงที่ผู้ใช้ส่งมาอีกรอบ) — ไฟล์ AMI ไม่ได้มีโครงสร้าง
# เดียวกันเป๊ะทุกไฟล์: หัวรายงานอยู่ Sheet1 (ป้ายไม่มีคำว่า "ไฟฟ้า" ต่อท้าย ต่างจากไฟล์แบบแรก)
# ตารางข้อมูลราย 15 นาทีอยู่ Sheet2 ต่างหาก (แต่ละ Rate มี 2 คอลัมน์ซ้ำค่ากัน) และปีในคอลัมน์
# เวลาเป็น ค.ศ. อยู่แล้ว (ไม่ใช่ พ.ศ. เหมือนไฟล์แบบแรก)
def _write_synthetic_multisheet_ami_workbook(path, account_no="0199000001", company_name="บริษัท หลายชีต จำกัด"):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A2"] = "รายงานข้อมูลกิโลวัตต์แบบช่วงเวลา"
    ws1["A3"] = "[ระหว่างวันที่ : 01 สิงหาคม 2568 - 31 สิงหาคม 2568]"
    ws1["A4"] = "บัญชีผู้ใช้ไฟ :"
    ws1["B4"] = account_no
    ws1["C4"] = "ชื่อผู้ใช้ไฟ :"
    ws1["D4"] = company_name
    ws1["A5"] = "หมายเลขมิเตอร์ :"
    ws1["B5"] = "1234567"
    ws1["C5"] = "Tariff :"
    ws1["D5"] = "TOU"
    ws1["A6"] = "CT Ratio :"
    ws1["B6"] = "100:5 A."
    ws1["C6"] = "VT Ratio :"
    ws1["D6"] = "115000:115 V."

    ws2 = wb.create_sheet("Sheet2")
    ws2["B2"] = "RATE A"
    ws2["C2"] = "RATE A"
    ws2["D2"] = "RATE B"
    ws2["E2"] = "RATE B"
    ws2["F2"] = "RATE C"
    ws2["G2"] = "RATE C"
    ws2["A3"] = "01/08/2025 00.15"
    ws2["D3"] = "3580.000"
    ws2["E3"] = "3580.000"
    ws2["A4"] = "01/08/2025 09.30"
    ws2["B4"] = "40.000"
    ws2["C4"] = "40.000"
    ws2["A5"] = "01/08/2025 22.15"
    ws2["F5"] = "15.000"
    ws2["G5"] = "15.000"
    ws2["A6"] = "กิโลวัตต์ต่ำสุด"
    ws2["B6"] = "20.000"

    ws3 = wb.create_sheet("Sheet3")
    ws3["A1"] = "***ค่าที่แสดงอาจไม่ตรงกับใบแจ้งค่าไฟฟ้า..."
    ws3["A3"] = "พิมพ์โดย : ทดสอบ"

    wb.save(path)


def test_parse_ami_interval_report_finds_rate_table_on_a_different_sheet_than_header(tmp_path):
    path = tmp_path / "multisheet.xlsx"
    _write_synthetic_multisheet_ami_workbook(path)

    readings = parse_ami_interval_report(path)

    # ปีในไฟล์นี้เป็น ค.ศ. อยู่แล้ว (ไม่ใช่ พ.ศ.) ต้องไม่ถูกลบ 543 ซ้ำอีก (จะกลายเป็นปี 1482 ผิด)
    assert IntervalReading(timestamp="01/08/2025 00.15", period="OP", kwh=3580.0) in readings
    assert IntervalReading(timestamp="01/08/2025 09.30", period="P", kwh=40.0) in readings
    assert IntervalReading(timestamp="01/08/2025 22.15", period="H", kwh=15.0) in readings
    assert len(readings) == 3


def test_parse_ami_report_header_finds_labels_when_not_on_first_sheet(tmp_path):
    path = tmp_path / "multisheet.xlsx"
    _write_synthetic_multisheet_ami_workbook(path, account_no="0277777777", company_name="บริษัท หลายชีต จำกัด")

    info = parse_ami_report_header(path)

    assert info["บัญชีผู้ใช้ไฟ"] == "0277777777"
    assert info["ชื่อผู้ใช้ไฟ"] == "บริษัท หลายชีต จำกัด"


# ── "Custom kW Report" — ไม่มีคอลัมน์ Rate A/B/C เลย (ยืนยันจากไฟล์จริงที่ผู้ใช้ส่งมา) ──


def _write_synthetic_custom_kw_report(path, account_no="0299999999", meter_no="1112223"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["B3"] = "Provincial Electricity Authority Advanced Metering Infrastructure (AMI) "
    ws["B6"] = "Custom kW Report"
    ws["B7"] = "Between 01 July 2025 - 31 July 2025"
    ws["B9"] = "Contact Account : "
    ws["C9"] = account_no
    ws["B10"] = "Meter No. : "
    ws["C10"] = meter_no

    ws["A28"] = "Time"
    ws["B28"] = "kW"
    # วันอังคารที่ 01/07/2025 09:30 -> วันทำการ ช่วง Peak (09-22)
    ws["A29"] = "01/07/2025 09.30"
    ws["B29"] = 100
    # เวลา 23:00 วันเดียวกัน -> วันทำการ ช่วง Off-Peak (22:00-09:00)
    ws["A30"] = "01/07/2025 23.00"
    ws["B30"] = 40
    # "24.00" ของวันที่ 01/07/2025 หมายถึงเที่ยงคืนของวันถัดไป (02/07/2025 00:00) ไม่ใช่ของวันเดิม
    ws["A31"] = "01/07/2025 24.00"
    ws["B31"] = 60

    wb.save(path)


def test_is_ami_xlsx_true_for_custom_kw_report(tmp_path):
    path = tmp_path / "kw_report.xlsx"
    _write_synthetic_custom_kw_report(path)
    assert is_ami_xlsx(path) is True


def test_parse_ami_report_header_reads_english_labels_from_custom_kw_report(tmp_path):
    path = tmp_path / "kw_report.xlsx"
    _write_synthetic_custom_kw_report(path, account_no="0288888888", meter_no="5556667")

    info = parse_ami_report_header(path)

    # ป้ายภาษาอังกฤษ ("Contact Account :", "Meter No. :") ต้องถูกแปลงให้ตรงกับ key มาตรฐาน
    # เดียวกับไฟล์อื่นๆ ไม่งั้นโค้ดฝั่งเรียกใช้ (amr_import.py) จะหาบัญชีผู้ใช้ไฟไม่เจอ
    assert info["บัญชีผู้ใช้ไฟ"] == "0288888888"
    assert info["หมายเลขมิเตอร์"] == "5556667"
    # ไฟล์รูปแบบนี้ไม่มีชื่อบริษัทให้เลย
    assert "ชื่อผู้ใช้ไฟ" not in info


def test_parse_ami_interval_report_derives_tou_period_from_timestamp_when_no_rate_columns(tmp_path):
    path = tmp_path / "kw_report.xlsx"
    _write_synthetic_custom_kw_report(path)

    readings = parse_ami_interval_report(path)

    # ไม่มีคอลัมน์ Rate A/B/C ให้เลย ต้องคำนวณช่วง P/OP/H เองจาก timestamp แล้วแปลง kW เป็น kWh
    # ของช่วง 15 นาที (kW x 0.25)
    assert IntervalReading(timestamp="01/07/2025 09.30", period="P", kwh=25.0) in readings
    assert IntervalReading(timestamp="01/07/2025 23.00", period="OP", kwh=10.0) in readings


def test_parse_ami_interval_report_handles_24_00_as_start_of_next_day(tmp_path):
    path = tmp_path / "kw_report.xlsx"
    _write_synthetic_custom_kw_report(path)

    readings = parse_ami_interval_report(path)

    # "01/07/2025 24.00" ต้องถูกเลื่อนเป็น "02/07/2025 00.00" (datetime.strptime ปกติ parse
    # ชั่วโมง 24 ไม่ได้ จะ error ถ้าไม่จัดการเป็นพิเศษ)
    assert IntervalReading(timestamp="02/07/2025 00.00", period="OP", kwh=15.0) in readings
    assert not any(r.timestamp.startswith("01/07/2025 24") for r in readings)
